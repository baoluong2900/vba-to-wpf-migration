"""
dd_common.py
------------
Shared fill engine of the four DD (Detail Design) Excel scripts:
``fill_dd_excel_frm.py`` (form), ``fill_dd_excel_report.py`` (report),
``fill_dd_excel_macro.py`` (macro) and ``fill_dd_excel_module.py`` (module).

Each script keeps its own ``fill_excel`` orchestration and CLI; everything they
had copied from each other lives here once:

- sheet-name patterns, renaming, ordering and the detail / query dispatchers
  (strict ``<id>_NNSelQueryMM`` plus DD_Report compact ``<id>_SelQueryMM`` names)
- 概要仕様書 (overview), 詳細仕様書 (``steps[]`` + embedded ``query_mappings`` tables)
- SelQuery / UpdQuery sheets: column grids, stacked SQL clause blocks, ``raw_sql``
  (used when Layer 2 leaves ``sql.select`` / ``sql.from`` empty), and the
  DD_Report compact SelQuery layout

Cell helpers (``put``, ``expand_table``, label lookup, change history, ...) come
from ``fill_bd_excel``.
"""
from __future__ import annotations

import copy
import re
import unicodedata
import warnings
from typing import Any

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from excel_translator import translate_label
from fill_bd_excel import (
    EXCEL_DETAIL_QUERY_EMBED_TABLE_FONT_SZ,
    EXCEL_DOCUMENT_FONT_NAME,
    apply_excel_document_font,
    clear_block,
    clear_row_cells,
    expand_table,
    fill_branded_header_fields,
    find_row_by_cell_text,
    purge_stale_data_rows,
    put,
    set_cell,
    template_span,
    write_text,
)

SHEET_OVERVIEW = "概要仕様書"


# ─────────────────────────────────────────────
# Sheet-name patterns
# ─────────────────────────────────────────────

# Templated names: 詳細仕様書_<prefix>_NN / <prefix>_NNSelQueryMM (screen family),
# 詳細仕様書_<id> / <id>_SelQueryMM (DD_Report). The prefix may contain underscores
# (e.g. "F_Login"); `.+` is greedy so the final digits anchor the split.
_RE_DETAIL   = re.compile(r"^詳細仕様書_(?P<prefix>.+)_(?P<nn>\d+)$")
_RE_DETAIL_LOOSE = re.compile(r"^詳細仕様書_(?P<id>.+)$")
_RE_SELQUERY = re.compile(r"^(?P<prefix>.+)_(?P<nn>\d+)SelQuery(?P<mm>\d+)$")
_RE_SELQUERY_COMPACT = re.compile(r"^(?P<prefix>.+)_SelQuery(?P<mm>\d+)$")
_RE_UPDQUERY = re.compile(r"^(?P<prefix>.+)_(?P<nn>\d+)UpdQuery(?P<mm>\d+)$")
_RE_UPDQUERY_COMPACT = re.compile(r"^(?P<prefix>.+)_UpdQuery(?P<mm>\d+)$")
_RE_INSQUERY = re.compile(r"^(?P<prefix>.+)_(?P<nn>\d+)InsQuery(?P<mm>\d+)$")
_RE_INSQUERY_COMPACT = re.compile(r"^(?P<prefix>.+)_InsQuery(?P<mm>\d+)$")
# Module-centric names used by form JSON patterns (no _NN token).
_RE_DETAIL_MODULE = re.compile(r"^詳細仕様書_(?P<module>[A-Za-z]\d{3})$")
_RE_QUERY_MODULE = re.compile(
    r"^(?P<module>[A-Za-z]\d{3})_(?P<kind>SelQuery|UpdQuery|InsQuery)(?P<mm>\d+)$",
)


def _safe_title(name: str) -> str:
    """Excel sheet titles are capped at 31 characters."""
    return name[:31]


def _safe_title_preserve_suffix(prefix: str, mid: str, suffix: str) -> str:
    """Trim middle text but keep suffix token (e.g. ``_01`` / ``SelQuery01``)."""
    max_len = 31
    title = f"{prefix}{mid}{suffix}"
    if len(title) <= max_len:
        return title
    keep = max_len - len(prefix) - len(suffix)
    if keep <= 0:
        return (prefix + suffix)[:max_len]
    return f"{prefix}{mid[:keep]}{suffix}"


def _prefix_matches(expected: str, actual: str) -> bool:
    """True when ``actual`` equals expected prefix or its truncated form."""
    return actual == expected or expected.startswith(actual)


def _match_query_sheet_name(name: str, *, kind: str) -> re.Match[str] | None:
    """Return a regex match for Sel / Upd / Ins query sheet titles."""
    if kind == "sel":
        m = _RE_SELQUERY.match(name)
        if m:
            return m
        return _RE_SELQUERY_COMPACT.match(name)
    m = _RE_UPDQUERY.match(name)
    if m:
        return m
    m = _RE_UPDQUERY_COMPACT.match(name)
    if m:
        return m
    m = _RE_INSQUERY.match(name)
    if m:
        return m
    return _RE_INSQUERY_COMPACT.match(name)


def detect_template_prefix(sheet_names: list[str]) -> str | None:
    """Detect the ``FD010M`` / ``RS020``-style prefix shared by detail/query sheets.

    Scans detail/query sheet names and returns the most common prefix.
    """
    counts: dict[str, int] = {}
    for n in sheet_names:
        m = _RE_DETAIL.match(n)
        if m:
            counts[m.group("prefix")] = counts.get(m.group("prefix"), 0) + 1
            continue
        m = _RE_DETAIL_LOOSE.match(n)
        if m and not _RE_DETAIL.match(n):
            rid = m.group("id").strip()
            if rid:
                counts[rid] = counts.get(rid, 0) + 5
            continue
        for rx in (_RE_SELQUERY, _RE_SELQUERY_COMPACT, _RE_UPDQUERY,
                   _RE_UPDQUERY_COMPACT, _RE_INSQUERY, _RE_INSQUERY_COMPACT):
            m = rx.match(n)
            if m:
                p = m.group("prefix").strip()
                if p:
                    counts[p] = counts.get(p, 0) + 1
                break
    if not counts:
        return None
    return max(counts.items(), key=lambda kv: kv[1])[0]


def _copy_worksheet_safely(wb, template_sheet):
    """Copy worksheet while suppressing transient >31-char title warning."""
    with warnings.catch_warnings():
        warnings.filterwarnings(
            "ignore",
            message=r"Title is more than 31 characters.*",
            category=UserWarning,
        )
        return wb.copy_worksheet(template_sheet)


# ─────────────────────────────────────────────
# Cell / border helpers
# ─────────────────────────────────────────────


def _unmerge_cells_touching(
    ws: Worksheet, min_row: int, max_row: int, min_col: int, max_col: int,
) -> None:
    """Drop merged ranges overlapping the rectangle so new merges can be applied."""
    to_drop: list[str] = []
    for rng in list(ws.merged_cells.ranges):
        if rng.max_row < min_row or rng.min_row > max_row:
            continue
        if rng.max_col < min_col or rng.min_col > max_col:
            continue
        to_drop.append(str(rng))
    for addr in to_drop:
        ws.unmerge_cells(addr)


def _apply_row_box_border(
    ws: Worksheet,
    row: int,
    *,
    col_start: int,
    col_end: int,
    side: Side,
) -> None:
    """Apply a boxed thin border across a full row span."""
    for c in range(col_start, col_end + 1):
        left = side if c == col_start else None
        right = side if c == col_end else None
        ws.cell(row, c).border = Border(
            left=left,
            right=right,
            top=side,
            bottom=side,
        )


def _apply_row_outer_border(
    ws: Worksheet,
    row: int,
    *,
    col_start: int,
    col_end: int,
    side: Side,
    draw_top: bool,
    draw_bottom: bool,
) -> None:
    """Apply only outer frame edges for a row within a multi-row merged block."""
    for c in range(col_start, col_end + 1):
        left = side if c == col_start else None
        right = side if c == col_end else None
        top = side if draw_top else None
        bottom = side if draw_bottom else None
        ws.cell(row, c).border = Border(
            left=left,
            right=right,
            top=top,
            bottom=bottom,
        )


def _duplicate_full_row_styles(
    ws: Worksheet,
    src_row: int,
    dst_row: int,
    max_col: int,
) -> None:
    h_src = ws.row_dimensions[src_row].height
    if h_src is not None:
        ws.row_dimensions[dst_row].height = h_src
    for c in range(1, max_col + 1):
        ws.cell(dst_row, c)._style = copy.copy(ws.cell(src_row, c)._style)


# ─────────────────────────────────────────────
# Stacked SQL blocks + 結果確認 band
# ─────────────────────────────────────────────


def apply_stacked_sql_blocks(
    ws: Worksheet,
    start_row: int,
    blocks: list[tuple[str, str]],
    *,
    col_start: int = 1,
    col_end: int = 39,
    body_col_start: int | None = None,
) -> int | None:
    """Write stacked SQL clause blocks: green title row + wrapped body row each.

    Returns the last worksheet row used for a **body** (white) line, or ``None``
    if nothing was written.
    """
    if col_start < 1 or col_end < col_start:
        return None
    body_col = body_col_start if body_col_start is not None else col_start
    if body_col < col_start or body_col > col_end:
        body_col = col_start
    hdr_fill = PatternFill(fill_type="solid", fgColor="C6EFCE")
    body_fill = PatternFill(fill_type="solid", fgColor="FFFFFF")
    hdr_font = Font(name=EXCEL_DOCUMENT_FONT_NAME, bold=True)
    # indent=0: Excel cell indent pushes body text inward vs header/data rows.
    body_align = Alignment(
        wrap_text=True, vertical="top", horizontal="left", indent=0,
    )
    hdr_align = Alignment(horizontal="left", vertical="center")
    thin_side = Side(style="thin", color="000000")
    row_border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side,
    )
    c1_letter = get_column_letter(col_start)
    c2_letter = get_column_letter(col_end)
    body_c1_letter = get_column_letter(body_col)
    r = start_row
    last_body_row: int | None = None
    for title, body in blocks:
        title_s = str(title or "").strip()
        body_s = str(body or "")
        body_s = body_s.replace("\\r\\n", "\n").replace("\\n", "\n").replace("\r\n", "\n").strip()
        if not title_s and not body_s:
            continue
        body_lines = [ln for ln in body_s.split("\n")] if body_s else [""]
        if not body_lines:
            body_lines = [""]
        end_row = r + len(body_lines)
        _unmerge_cells_touching(ws, r, end_row, col_start, col_end)
        merge_addr = f"{c1_letter}{r}:{c2_letter}{r}"
        ws.merge_cells(merge_addr)
        top = ws.cell(r, col_start)
        top.value = title_s
        top.font = hdr_font
        top.fill = hdr_fill
        top.alignment = hdr_align
        top.border = row_border
        apply_excel_document_font(top)
        _apply_row_box_border(
            ws, r, col_start=col_start, col_end=col_end, side=thin_side,
        )
        for i, ln in enumerate(body_lines):
            r += 1
            merge_addr2 = f"{body_c1_letter}{r}:{c2_letter}{r}"
            ws.merge_cells(merge_addr2)
            if body_col > col_start:
                lead = ws.cell(r, col_start)
                lead.value = None
                lead.fill = body_fill
                apply_excel_document_font(lead)
            cell = ws.cell(r, body_col)
            write_text(cell, ln[:32000] if ln else None)
            cell.fill = body_fill
            cell.alignment = body_align
            cell.border = Border()
            _apply_row_outer_border(
                ws,
                r,
                col_start=col_start,
                col_end=col_end,
                side=thin_side,
                draw_top=(i == 0),
                draw_bottom=(i == len(body_lines) - 1),
            )
            apply_excel_document_font(cell)
            last_body_row = r
        r += 1
    return last_body_row


_RESULT_CONFIRM_LABEL = "結果確認"


def _resolve_result_confirm_body_band(
    ws: Worksheet, body_style_row: int,
) -> tuple[int, int, bool]:
    """Return (col_lo, col_hi, merge_each_row) for 結果確認 body cells on ``body_style_row``."""
    ao_idx = column_index_from_string("AO")
    for mr in ws.merged_cells.ranges:
        if mr.min_row <= body_style_row <= mr.max_row and mr.min_col <= ao_idx <= mr.max_col:
            if mr.min_row == mr.max_row == body_style_row:
                tl = ws.cell(mr.min_row, mr.min_col).value
                if str(tl or "").strip() == _RESULT_CONFIRM_LABEL:
                    break
                return mr.min_col, mr.max_col, True
            break
    return ao_idx, ao_idx + 1, False


def _stretch_sel_query_result_confirm_body(
    ws: Worksheet,
    *,
    body_style_row: int,
    sql_body_last_row: int | None,
) -> None:
    """Extend 結果確認 body styling (column AO band) through stacked SQL body rows."""
    if sql_body_last_row is None or body_style_row >= sql_body_last_row:
        return
    if find_row_by_cell_text(
        ws, "AO", _RESULT_CONFIRM_LABEL, row_start=1, row_end=120,
    ) is None:
        return
    col_lo, col_hi, merge_each_row = _resolve_result_confirm_body_band(
        ws, body_style_row,
    )
    rh = ws.row_dimensions[body_style_row].height
    for r in range(body_style_row + 1, sql_body_last_row + 1):
        _unmerge_cells_touching(ws, r, r, col_lo, col_hi)
        for c in range(col_lo, col_hi + 1):
            dst = ws.cell(r, c)
            dst.value = None
            dst._style = copy.copy(ws.cell(body_style_row, c)._style)
            apply_excel_document_font(dst)
        if rh is not None:
            ws.row_dimensions[r].height = rh
        if merge_each_row:
            ws.merge_cells(
                start_row=r, start_column=col_lo, end_row=r, end_column=col_hi,
            )


def find_sel_query_column_header_row(
    ws: Worksheet,
    *,
    after_row: int,
    scan_rows: int = 35,
) -> int | None:
    """First column-grid header row after ``SELECT句`` (``A`` matches No / variants)."""
    row_start = after_row + 1
    row_end = after_row + scan_rows
    for marker in ("No", "No.", "NO", "№", "Ｎｏ．"):
        r = find_row_by_cell_text(ws, "A", marker, row_start=row_start, row_end=row_end)
        if r is not None:
            return r
    return None


# ─────────────────────────────────────────────
# Common header (rows 1-5) shared by overview / detail / query sheets
# ─────────────────────────────────────────────


def fill_common_header(ws: Worksheet, dd: dict, *,
                       module_id: str | None = None,
                       module_name: str | None = None,
                       test_item_count: Any = None,
                       test_completed_count: Any = None) -> None:
    """Fill the 5-row banner present on overview / detail / query sheets.

    Cells (merge-aware; top-left coord used):
        E1  system_name        AD1 author           AK1 created_date   (by label)
        E2  subsystem_name     AD2 updated_by       AK2 updated_date
        E4  screen_id          R4  screen_name
        E5  module_id          R5  module_name
        AO2 test_item_count    AP2 test_completed_count   (detail sheets only)

    Not for 変更履歴: that sheet has no banner (its row 1-2 merges are the title
    and the table header).
    """
    fill_branded_header_fields(ws, dd)
    put(ws, "E2",  dd.get("subsystem_name"))
    put(ws, "AD2", dd.get("updated_by"))
    put(ws, "AK2", dd.get("updated_date"))
    put(ws, "E4",  dd.get("screen_id"))
    put(ws, "R4",  dd.get("screen_name"))
    if module_id is not None:
        put(ws, "E5", module_id)
    if module_name is not None:
        put(ws, "R5", module_name)
    if test_item_count is not None:
        put(ws, "AO2", test_item_count)
    if test_completed_count is not None:
        put(ws, "AP2", test_completed_count)


# ─────────────────────────────────────────────
# 概要仕様書 (Overview)
# ─────────────────────────────────────────────

OV_PARAM_MARKERS  = [("A", "パラメータ情報")]   # section header (merged A:AM)
OV_PARAM_FALLBACK = 15                           # first data row below header

# DD_Screen template (概要仕様書 row 14 headers); DD_Report passes its own map.
OV_PARAM_COL = {
    "no":             "A",   # merged A:B
    "io":             "C",   # merged C:D
    "item_name":      "E",   # merged E:M
    "symbol_name":    "N",   # merged N:V
    "data_type":      "W",   # merged W:X
    "total_digits":   "Y",   # merged Y:Z
    "decimal_digits": "AA",  # merged AA:AB
    "required":       "AC",  # merged AC:AD
    "notes":          "AE",  # merged AE:AM   (説明)
}


def fill_overview(ws: Worksheet, dd: dict, *,
                  param_cols: dict[str, str] | None = None) -> None:
    param_cols = param_cols or OV_PARAM_COL
    fill_common_header(ws, dd, module_id="概要", module_name=None)

    ov = dd.get("sheet_2_overview", {}) or {}

    # Description (single free-text cell under A5="概要" banner).
    # Keep the row-5 label intact; write description into the row below.
    desc = ov.get("description")
    if desc:
        put(ws, "B6", desc)

    param_hdr_row = None
    r = find_row_by_cell_text(ws, "A", "パラメータ情報")
    if r is not None:
        col_hdr = find_row_by_cell_text(ws, "A", "No",
                                        row_start=r + 1, row_end=r + 5)
        param_hdr_row = col_hdr if col_hdr is not None else (r + 1)
    if param_hdr_row is None:
        param_hdr_row = OV_PARAM_FALLBACK - 1
    template_row = param_hdr_row + 1

    params = ov.get("parameter_info", []) or []

    notes_banner_row = find_row_by_cell_text(
        ws, "A", "特記事項",
        row_start=template_row,
        row_end=min(template_row + 120, ws.max_row or 300),
    )
    if notes_banner_row is None:
        notes_banner_row = param_hdr_row + 2

    banner_style_row = notes_banner_row
    max_col = ws.max_column or 39
    banner_styles = [
        copy.copy(ws.cell(banner_style_row, c)._style)
        for c in range(1, max_col + 1)
    ]
    banner_height = ws.row_dimensions[banner_style_row].height

    gap = max(1, notes_banner_row - template_row)
    n_par = len(params)
    if n_par > gap:
        ws.insert_rows(notes_banner_row, amount=n_par - gap)
        notes_banner_row += n_par - gap

    cols = param_cols.values()
    for r in range(template_row, notes_banner_row):
        clear_row_cells(ws, r, cols)

    if params:
        rows = expand_table(ws, template_row, n_par)
        for rr in rows:
            if rr != template_row:
                _duplicate_full_row_styles(ws, template_row, rr, max_col=max_col)
        for item, rr in zip(params, rows):
            clear_row_cells(ws, rr, cols)
            for key, col in param_cols.items():
                if key in item:
                    put(ws, f"{col}{rr}", item[key])

    for rng in list(ws.merged_cells.ranges):
        if rng.min_row <= notes_banner_row <= rng.max_row:
            try:
                ws.unmerge_cells(str(rng))
            except (KeyError, ValueError):
                pass
    ws.merge_cells(
        start_row=notes_banner_row, start_column=1,
        end_row=notes_banner_row, end_column=min(39, max_col),
    )
    for c in range(1, max_col + 1):
        ws.cell(notes_banner_row, c)._style = copy.copy(banner_styles[c - 1])
    if banner_height is not None:
        ws.row_dimensions[notes_banner_row].height = banner_height
    top = ws.cell(notes_banner_row, 1).value
    if top in (None, "") or str(top).strip() == "":
        put(ws, f"A{notes_banner_row}", translate_label("特記事項"))

    notes = ov.get("special_notes")
    if notes:
        sn_row = notes_banner_row + 1
        put(ws, f"A{sn_row}", notes)


def has_meaningful_overview_payload(dd: dict) -> bool:
    """Return True if sheet_2_overview carries real overview content.

    For module/macro templates without 概要仕様書, many legacy JSON files keep
    only ``sheet_2_overview.details`` as a transport for detail specs; that
    should not trigger a warning.
    """
    ov = dd.get("sheet_2_overview")
    if not isinstance(ov, dict):
        return False
    if (ov.get("description") or "").strip():
        return True
    if (ov.get("special_notes") or "").strip():
        return True
    params = ov.get("parameter_info")
    if isinstance(params, list) and len(params) > 0:
        return True
    gsp = ov.get("grouping_sorting_pagebreak")
    if isinstance(gsp, dict):
        for k in ("sort_fields", "group_fields"):
            v = gsp.get(k)
            if isinstance(v, list) and len(v) > 0:
                return True
    return False


# ─────────────────────────────────────────────
# 詳細仕様書_* (per-module detail sheets)
# ─────────────────────────────────────────────


def _detail_sheets_in_order(wb, tmpl_prefix: str) -> list[str]:
    """Return detail sheet names (strict ``_NN`` or loose ``詳細仕様書_<id>``) in order."""
    strict: list[tuple[int, str]] = []
    loose: list[str] = []
    for n in wb.sheetnames:
        m = _RE_DETAIL.match(n)
        if m and _prefix_matches(tmpl_prefix, m.group("prefix")):
            strict.append((int(m.group("nn")), n))
            continue
        m = _RE_DETAIL_LOOSE.match(n)
        if m and not _RE_DETAIL.match(n):
            if _prefix_matches(tmpl_prefix, m.group("id").strip()):
                loose.append(n)
    strict.sort(key=lambda kv: kv[0])
    loose.sort()
    return [name for _, name in strict] + loose


def _resolve_detail_sheet_name(detail: dict, screen_id: str, fallback_idx: int) -> str:
    """Build target detail sheet name for sheet_2_detail items."""
    pattern = str(detail.get("sheet_name_pattern") or "").strip()
    if pattern:
        nn = str(detail.get("nn") or detail.get("module_no") or (fallback_idx + 1)).zfill(2)
        name = (pattern
                .replace("{screen_id}", screen_id)
                .replace("{FormID}", screen_id)
                .replace("{nn}", nn))
        return _safe_title(name)
    return _safe_title_preserve_suffix(
        "詳細仕様書_",
        screen_id,
        f"_{fallback_idx+1:02d}",
    )


def _detail_sheet_is_dd_report_process_band(ws: Worksheet) -> bool:
    """True for DD_Report detail: ``帳票ID`` row sits directly above ``入力データ欄``.

    On those sheets, ``E5``/``R5`` are inside the process-table header merges; writing
    module_id there overwrites ``入力データ欄`` / ``処理記述`` and breaks step fill.
    Screen DD uses an extra ``モジュールID`` row so ``E5``/``R5`` are safe.
    """
    r_ch = find_row_by_cell_text(ws, "A", "帳票ID", row_start=3, row_end=10)
    r_in = find_row_by_cell_text(ws, "A", "入力データ欄", row_start=5, row_end=15)
    if r_ch is None or r_in is None:
        return False
    return r_in == r_ch + 1


def fill_detail_sheet(ws: Worksheet, dd: dict, detail: dict, *,
                      annotate_events: bool = False) -> None:
    """Fill header + module metadata + ``steps[]`` on a single detail sheet.

    DD_Report detail templates omit the ``モジュールID`` row; module_id / module_name
    must not be written to ``E5``/``R5`` (those cells belong to the green banner).

    ``annotate_events`` (macro / module): surface ``event_trigger`` /
    ``event_content`` in AO7 when that cell is empty, so nothing is silently dropped.
    """
    report_band = _detail_sheet_is_dd_report_process_band(ws)
    fill_common_header(
        ws, dd,
        module_id=None if report_band else detail.get("module_id"),
        module_name=None if report_band else detail.get("module_name"),
        test_item_count=detail.get("test_item_count"),
        test_completed_count=detail.get("test_completed_count"),
    )

    if annotate_events:
        trig = (detail.get("event_trigger") or "").strip()
        cont = (detail.get("event_content") or "").strip()
        if (trig or cont) and ws.cell(7, 41).value in (None, ""):   # AO = col 41
            # Below the 結果確認 header (AO6); 入力データ欄/処理記述 labels stay intact.
            set_cell(ws, "AO7", " / ".join(s for s in (trig, cont) if s))

    _fill_detail_steps(ws, detail)


_DETAIL_SPEC_BODY_COLS: tuple[str, ...] = (
    tuple(chr(c) for c in range(ord("A"), ord("Z") + 1))
    + (
        "AA", "AB", "AC", "AD", "AE", "AF", "AG", "AH",
        "AI", "AJ", "AK", "AL", "AM",
    )
)

# Rows below the green 入力データ欄/処理記述/出力データ欄 header: first body row (style
# template) + one closing border row. Do not derive this from ws.max_row — extra
# rows below the sheet would shrink gap and skip inserts.
DETAIL_PROCESS_TABLE_STUB_ROWS = 2


def _detail_process_header_row(ws: Worksheet) -> int | None:
    """Locate the green banner row ``入力データ欄`` / ``処理記述`` / ``出力データ欄``."""
    r = find_row_by_cell_text(ws, "H", "処理記述", row_start=1, row_end=40)
    if r is not None:
        return r
    return find_row_by_cell_text(ws, "A", "入力データ欄", row_start=1, row_end=40)


# Detail「処理記述」: numbered outline (I..AF) + optional query_mappings tables.
DETAIL_SPEC_MIN_COL_IDX = 9
DETAIL_SPEC_MAX_COL_IDX = 32
DETAIL_SPEC_WRAP_UNITS_PER_COL = 5

_DETAIL_QUERY_PLACEHOLDER_RE = re.compile(r"\n?\s*\[([^\]\r\n]+)\]\s*\n?")
_DETAIL_QUERY_TABLE_SIDE = Side(style="thin", color="000000")
_DETAIL_QUERY_TABLE_SPAN_COLS = 13
_DETAIL_QUERY_TABLE_BAND1_COLS = 4
_DETAIL_QUERY_TABLE_BAND2_COLS = 2
_DETAIL_QUERY_TABLE_FILL_HDR = PatternFill(fill_type="solid", fgColor="C6EFCE")
_DETAIL_QUERY_TABLE_FILL_BODY = PatternFill(fill_type="solid", fgColor="FFFFFF")

_DETAIL_BODY_SPEC_LINE_NUM_RE = re.compile(
    r"^\s*(\d+(?:\.\d+)*)(?:\s*[\.\．]\s*|\s+|$)",
)


def _detail_body_spec_num_segments(line: str) -> int | None:
    s = line.strip()
    if not s:
        return None
    m = _DETAIL_BODY_SPEC_LINE_NUM_RE.match(s)
    if not m:
        return None
    return len(m.group(1).split("."))


def _detail_body_spec_line_column_letter(
    line: str, prev_indent: int,
) -> tuple[str, int]:
    segs = _detail_body_spec_num_segments(line)
    if segs is None:
        indent = prev_indent
    else:
        indent = max(0, segs - 2)
    col_idx = DETAIL_SPEC_MIN_COL_IDX + indent
    if col_idx > DETAIL_SPEC_MAX_COL_IDX:
        col_idx = DETAIL_SPEC_MAX_COL_IDX
    return get_column_letter(col_idx), indent


def _detail_body_char_display_units(ch: str) -> int:
    if not ch:
        return 0
    return 2 if unicodedata.east_asian_width(ch) in ("F", "W", "A") else 1


def _wrap_detail_body_line_to_width(
    text: str, *, start_col_idx: int,
) -> list[str]:
    span_cols = DETAIL_SPEC_MAX_COL_IDX - start_col_idx + 1
    if span_cols < 1:
        span_cols = 1
    max_units = max(8, span_cols * DETAIL_SPEC_WRAP_UNITS_PER_COL)
    chunks: list[str] = []
    buf: list[str] = []
    units = 0
    for ch in text:
        u = _detail_body_char_display_units(ch)
        if buf and units + u > max_units:
            chunks.append("".join(buf))
            buf = [ch]
            units = u
        else:
            buf.append(ch)
            units += u
    if buf:
        chunks.append("".join(buf))
    if not chunks:
        chunks = [""]
    return chunks


def _index_query_mappings(detail: dict) -> dict[str, dict]:
    raw = detail.get("query_mappings")
    if not isinstance(raw, list):
        return {}
    out: dict[str, dict] = {}
    for item in raw:
        if not isinstance(item, dict):
            continue
        sn = str(item.get("sheet_name") or "").strip()
        if sn:
            out[sn] = item
    return out


def _query_mapping_table_rows(mapping: dict) -> list[tuple[str, str, str]]:
    sheet = str(mapping.get("sheet_name") or "").strip()
    if not sheet:
        sheet = str(mapping.get("logic_query_name") or "").strip()
    rows: list[tuple[str, str, str]] = [
        (translate_label("Query名"), sheet, ""),
        (translate_label("パラメータ"), "I/O", translate_label("編集仕様")),
    ]
    params = mapping.get("parameters")
    if isinstance(params, list):
        for p in params:
            if not isinstance(p, dict):
                continue
            name = str(p.get("name") or "").strip()
            io = str(p.get("io") or "").strip()
            edit = str(p.get("mapping") or p.get("edit_spec") or "").strip()
            rows.append((name, io, edit))
    return rows


def _expand_detail_body_for_layout(
    body: str,
    query_maps: dict[str, dict],
) -> list[tuple[str, ...]]:
    rows_out: list[tuple[str, ...]] = []
    prev_indent = 0
    pos = 0
    for m in _DETAIL_QUERY_PLACEHOLDER_RE.finditer(body):
        if m.start() > pos:
            chunk = body[pos:m.start()]
            lines = chunk.splitlines()
            if not lines:
                lines = [chunk] if chunk else [""]
            for line in lines:
                col_letter, prev_indent = _detail_body_spec_line_column_letter(
                    line, prev_indent,
                )
                col_idx = column_index_from_string(col_letter)
                for piece in _wrap_detail_body_line_to_width(
                    str(line), start_col_idx=col_idx,
                ):
                    rows_out.append(("single", col_letter, piece))
        indent_for_table = prev_indent
        key = m.group(1).strip()
        q = query_maps.get(key)
        if q is None:
            rows_out.append(
                ("single", get_column_letter(DETAIL_SPEC_MIN_COL_IDX), f"[{key}]"),
            )
        else:
            rows_out.append(("blank",))
            for trip in _query_mapping_table_rows(q):
                rows_out.append(("triple", indent_for_table, *trip))
            rows_out.append(("blank",))
        pos = m.end()
    if pos < len(body):
        chunk = body[pos:]
        lines = chunk.splitlines()
        if not lines:
            lines = [chunk] if chunk else [""]
        for line in lines:
            col_letter, prev_indent = _detail_body_spec_line_column_letter(
                line, prev_indent,
            )
            col_idx = column_index_from_string(col_letter)
            for piece in _wrap_detail_body_line_to_width(
                str(line), start_col_idx=col_idx,
            ):
                rows_out.append(("single", col_letter, piece))
    return rows_out


def _detail_steps_body_rows_needed(steps: list[Any], detail: dict) -> int:
    total = 0
    qmaps = _index_query_mappings(detail)
    for s in steps:
        if not isinstance(s, dict):
            continue
        body = str(s.get("detail") or "")
        expanded = _expand_detail_body_for_layout(body, qmaps)
        total += 1 + len(expanded) + 1
    return total


def _detail_query_table_band_edges(*, indent_shift: int = 0) -> tuple[int, int, int, int]:
    hi_full = DETAIL_SPEC_MAX_COL_IDX
    lo = DETAIL_SPEC_MIN_COL_IDX + max(0, indent_shift)
    if hi_full - lo + 1 < 10:
        lo = DETAIL_SPEC_MIN_COL_IDX
    span = min(_DETAIL_QUERY_TABLE_SPAN_COLS, hi_full - lo + 1)
    hi = lo + span - 1
    w1 = min(
        _DETAIL_QUERY_TABLE_BAND1_COLS,
        max(3, span - _DETAIL_QUERY_TABLE_BAND2_COLS - 2),
    )
    w2 = min(_DETAIL_QUERY_TABLE_BAND2_COLS, max(2, span - w1 - 2))
    e1 = lo + w1 - 1
    e2 = e1 + w2
    if e2 >= hi:
        e2 = hi - 1
    if e2 <= e1:
        e2 = min(e1 + 1, hi - 1)
    return lo, e1, e2, hi


def _outline_detail_query_band_border(ws: Worksheet, row: int, bc: int, ec: int) -> None:
    thin = _DETAIL_QUERY_TABLE_SIDE
    for c in range(bc, ec + 1):
        cell = ws.cell(row=row, column=c)
        cell.border = Border(
            left=thin if c == bc else None,
            right=thin if c == ec else None,
            top=thin,
            bottom=thin,
        )


def _detail_query_escape_plain_text(val: str) -> str | None:
    """Trimmed cell text (None when empty); written with ``write_text``."""
    s = str(val or "").strip()
    if not s:
        return None
    return s[:32000]


def _style_detail_query_merged_cell(
    cell,
    *,
    fill: PatternFill,
    bold: bool,
) -> None:
    cell.fill = fill
    cell.font = Font(
        name=EXCEL_DOCUMENT_FONT_NAME,
        sz=EXCEL_DETAIL_QUERY_EMBED_TABLE_FONT_SZ,
        bold=bold,
    )
    cell.alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")


def _write_detail_query_table_row(
    ws: Worksheet,
    row: int,
    *,
    row_index_in_table: int,
    indent_shift: int = 0,
    v1: str,
    v2: str,
    v3: str,
) -> None:
    lo, e1, e2, hi = _detail_query_table_band_edges(indent_shift=indent_shift)
    _unmerge_cells_touching(ws, row, row, lo, hi)
    for cc in range(lo, hi + 1):
        ws.cell(row=row, column=cc).value = None

    if row_index_in_table == 0:
        label = _detail_query_escape_plain_text(v1) or translate_label("Query名")
        qval = _detail_query_escape_plain_text(v2) or _detail_query_escape_plain_text(v3)

        ws.merge_cells(start_row=row, start_column=lo, end_row=row, end_column=e1)
        c_label = ws.cell(row=row, column=lo)
        write_text(c_label, label)
        _style_detail_query_merged_cell(
            c_label, fill=_DETAIL_QUERY_TABLE_FILL_HDR, bold=False,
        )
        _outline_detail_query_band_border(ws, row, lo, e1)

        ws.merge_cells(start_row=row, start_column=e1 + 1, end_row=row, end_column=hi)
        c_val = ws.cell(row=row, column=e1 + 1)
        write_text(c_val, qval)
        _style_detail_query_merged_cell(
            c_val, fill=_DETAIL_QUERY_TABLE_FILL_BODY, bold=False,
        )
        _outline_detail_query_band_border(ws, row, e1 + 1, hi)
        return

    bands = (
        (lo, e1),
        (e1 + 1, e2),
        (e2 + 1, hi),
    )
    vals = (
        _detail_query_escape_plain_text(v1),
        _detail_query_escape_plain_text(v2),
        _detail_query_escape_plain_text(v3),
    )
    is_hdr_row = row_index_in_table == 1
    hdr_fill = _DETAIL_QUERY_TABLE_FILL_HDR
    body_fill = _DETAIL_QUERY_TABLE_FILL_BODY
    for (bc, ec), val in zip(bands, vals):
        ws.merge_cells(start_row=row, start_column=bc, end_row=row, end_column=ec)
        top = ws.cell(row=row, column=bc)
        write_text(top, val)
        _style_detail_query_merged_cell(
            top,
            fill=hdr_fill if is_hdr_row else body_fill,
            bold=False,
        )
        _outline_detail_query_band_border(ws, row, bc, ec)


def _paint_detail_body_expanded(
    ws: Worksheet,
    detail_row_start: int,
    expanded: list[tuple[str, ...]],
) -> None:
    prev_was_triple = False
    triple_run_idx = 0
    for j, row_spec in enumerate(expanded):
        rnum = detail_row_start + j
        if row_spec[0] == "single":
            prev_was_triple = False
            _, col_letter, text_line = row_spec
            put(ws, f"{col_letter}{rnum}", text_line)
        elif row_spec[0] == "blank":
            prev_was_triple = False
        else:
            if not prev_was_triple:
                triple_run_idx = 0
            else:
                triple_run_idx += 1
            prev_was_triple = True
            _, indent_shift, vi, vj, vk = row_spec
            _write_detail_query_table_row(
                ws,
                rnum,
                row_index_in_table=triple_run_idx,
                indent_shift=int(indent_shift),
                v1=str(vi),
                v2=str(vj),
                v3=str(vk),
            )


def _ensure_detail_sheet_body_capacity(
    ws: Worksheet,
    hdr_row: int,
    min_body_rows: int,
    max_col: int,
) -> None:
    """Insert rows beneath the banner so ``min_body_rows`` fit before the footer line.

    Templates ship a short stub (often 2 rows: interior + bottom border row).
    We insert gap rows immediately above that bottom stub so borders stay closed.
    """
    if min_body_rows < 1:
        return
    stub_rows = DETAIL_PROCESS_TABLE_STUB_ROWS
    gap = min_body_rows - stub_rows
    if gap <= 0:
        return
    insert_before = hdr_row + stub_rows
    style_src = hdr_row + 1
    ws.insert_rows(insert_before, amount=gap)
    for off in range(gap):
        _duplicate_full_row_styles(
            ws, style_src, insert_before + off, max_col=max_col,
        )


def _fill_detail_steps(ws: Worksheet, detail: dict) -> None:
    """Fill every JSON ``steps[]`` entry; pre-expand stub rows then insert per-wrap lines."""
    steps = detail.get("steps")
    if not isinstance(steps, list):
        return

    hdr_row = _detail_process_header_row(ws)
    if hdr_row is None:
        return

    max_col = ws.max_column or 39

    rows_needed = _detail_steps_body_rows_needed(steps, detail)
    cols = list(_DETAIL_SPEC_BODY_COLS)
    body_top = hdr_row + 1
    if rows_needed >= 1:
        _ensure_detail_sheet_body_capacity(ws, hdr_row, rows_needed, max_col=max_col)

    for r in range(body_top, ws.max_row + 1):
        clear_row_cells(ws, r, cols)

    if rows_needed < 1:
        return

    qmaps = _index_query_mappings(detail)
    cur = body_top
    for i, s in enumerate(steps):
        if not isinstance(s, dict):
            continue
        title = s.get("step") or f"{i+1}. {translate_label('処理')}"
        body = s.get("detail") or ""
        put(ws, f"H{cur}", title)
        expanded = _expand_detail_body_for_layout(str(body), qmaps)

        detail_row_start = cur + 1
        _paint_detail_body_expanded(ws, detail_row_start, expanded)

        cur = detail_row_start + len(expanded) + 1


# ─────────────────────────────────────────────
# SQL text helpers
# ─────────────────────────────────────────────

# Column mapping for SELECT query data rows (row 11+).
SEL_COL = {
    "no":           "A",   # merged A
    "table_name":   "B",   # merged B:H
    "column_name":  "I",   # merged I:O
    "display_name": "P",   # merged P:V
    "remark":       "W",   # merged W:AM
}

# Stacked SELECT/FROM/WHERE blocks under both Sel and Upd grids (same width as Sheet4 Sel).
_DD_STACKED_SQL_COL_START = 1
_DD_STACKED_SQL_COL_END = 39
_DD_STACKED_SQL_BODY_COL_START = 2

# Column mapping for UPDATE query data rows (row 11+); see _resolve_upd_col_map.
UPD_COL = {
    "no":             "A",
    "column_name":    "B",   # merged B:J
    "source_area":    "K",   # merged K:S  (転送元 領域名)
    "source_field":   "T",   # merged T:AB (転送元 項目名／固定値)
    "remark":         "AC",  # merged AC:AM
}

_UPD_FIELD_KEYS: tuple[str, ...] = (
    "no", "column_name", "source_area", "source_field", "remark",
)


def _normalize_upd_header_label(val: Any) -> str:
    return "".join(ch for ch in str(val or "") if ch not in " \t\n\r\u3000")


def _row_horizontal_segments(
    ws: Worksheet, r: int, *, max_col: int = 45,
) -> list[tuple[int, int, Any]]:
    """Left-to-right header segments on a single row (honour horizontal merges)."""
    segments: list[tuple[int, int, Any]] = []
    c = 1
    while c <= max_col:
        hit = None
        for mr in ws.merged_cells.ranges:
            if mr.min_row == mr.max_row == r and mr.min_col <= c <= mr.max_col:
                hit = mr
                break
        if hit is not None:
            v = ws.cell(hit.min_row, hit.min_col).value
            segments.append((hit.min_col, hit.max_col, v))
            c = hit.max_col + 1
        else:
            segments.append((c, c, ws.cell(r, c).value))
            c += 1
    return segments


def _resolve_upd_col_map(ws: Worksheet, hdr_row: int) -> dict[str, str]:
    """Map logical fields to column letters from the template header row.

    Supports:
        - compact 5-column layout (No / カラム名 / 領域 / 項目 / 備考)
        - wide merged layout (same labels, ``UPD_COL``-style spans)
    """
    fallback = dict(UPD_COL)
    segs = _row_horizontal_segments(ws, hdr_row)
    found: dict[str, str | None] = {
        "no": None,
        "column_name": None,
        "source_area": None,
        "source_field": None,
        "remark": None,
    }
    for lo, _hi, raw in segs:
        label_norm = _normalize_upd_header_label(raw)
        if not label_norm and raw is None:
            continue
        le = get_column_letter(lo)
        raw_strip = str(raw or "").strip()
        rus = raw_strip.replace(".", "").upper()

        # Japanese labels, then the English template ones ("Column Name", ...).
        if rus in ("NO", "NO.", "NO．", "ＮＯ", "№"):
            slot = "no"
        elif "カラム名" in label_norm or "ColumnName" in label_norm:
            slot = "column_name"
        elif "備考" in label_norm or "Notes" in label_norm:
            slot = "remark"
        elif "領域名" in label_norm or "Region" in label_norm:
            slot = "source_area"
        elif any(t in label_norm for t in ("項目", "固定値", "固定", "Field", "Fixed")):
            slot = "source_field"
        elif (
            "転送元" in label_norm
            and "項目" not in label_norm
            and "固定" not in label_norm
        ):
            slot = "source_area"
        else:
            continue
        if found[slot] is None:
            found[slot] = le
    if all(found[k] for k in found):
        return {k: str(found[k]) for k in found}
    return fallback


# Placeholder values Layer 2 writes for an empty SQL clause.
_SQL_EMPTY_VALUES = ("", "なし", "N/A")


def _sql_clause(value: Any) -> str:
    """Clause text, or "" for an empty / placeholder clause ("", "なし", "N/A").

    Lists (``["A.ID", ...]`` / ``[{"expression": ...}, ...]``) give one item per line.
    """
    if value is None:
        return ""
    if isinstance(value, list):
        parts: list[str] = []
        for item in value:
            if isinstance(item, dict):
                item = (item.get("expression") or item.get("clause")
                        or item.get("text") or item.get("column"))
            text = _sql_clause(item)
            if text:
                parts.append(text)
        return "\n".join(parts)
    text = str(value).strip()
    return "" if text.upper() in _SQL_EMPTY_VALUES else text


def _format_sql_joins(sql: dict[str, Any]) -> str:
    j = sql.get("join")
    if not j:
        return ""
    if isinstance(j, list):
        lines: list[str] = []
        for x in j:
            if isinstance(x, dict):
                ln = str(
                    x.get("clause")
                    or x.get("text")
                    or x.get("join")
                    or x.get("expression")
                    or ""
                ).strip()
                # Standardized join object: {"type": "...", "table": "...", "on": "..."}
                if not ln:
                    jt = str(x.get("type") or "JOIN").strip()
                    tbl = str(x.get("table") or "").strip()
                    on = str(x.get("on") or "").strip()
                    if tbl and on:
                        ln = f"{jt} {tbl} ON {on}"
                    elif tbl:
                        ln = f"{jt} {tbl}"
            else:
                ln = str(x).strip()
            ln = _sql_clause(ln)
            if ln:
                lines.append(ln)
        return "\n".join(lines)
    return _sql_clause(j)


def _where_rows_from_sql(sql: dict[str, Any]) -> list[Any]:
    raw = sql.get("where")
    if isinstance(raw, list):
        return [x for x in raw if x is not None]
    if raw is not None and str(raw).strip():
        return [{"condition": raw, "notes": ""}]
    return []


def _where_condition_for_cell(item: Any) -> tuple[str, str]:
    """Return (condition text for column C, optional notes for column W)."""
    if isinstance(item, dict):
        c = (item.get("condition") or item.get("expression") or "").strip()
        notes = (item.get("notes") or item.get("remark") or "").strip()
    else:
        c = str(item).strip()
        notes = ""
    return c, notes


def _where_lines(sql: dict[str, Any]) -> list[str]:
    """WHERE conditions (placeholders dropped), each followed by its ``# notes``."""
    lines: list[str] = []
    for witem in _where_rows_from_sql(sql):
        cond, notes = _where_condition_for_cell(witem)
        cond = _sql_clause(cond)
        if cond:
            lines.append(cond + (f"  # {notes}" if notes else ""))
    return lines


def _normalize_sql_newlines(s: str) -> str:
    return (
        str(s)
        .replace("\\r\\n", "\n")
        .replace("\\n", "\n")
        .replace("\r\n", "\n")
        .strip()
    )


def _build_sel_sql_section_blocks(sql: dict[str, Any]) -> list[tuple[str, str]]:
    """Ordered (section title, SQL text) pairs for stacked SQL blocks; empty /
    placeholder clauses ("", "なし", "N/A") are left out."""
    blocks: list[tuple[str, str]] = []
    for title, key in (("SELECT", "select"), ("FROM", "from")):
        text = _sql_clause(sql.get(key))
        if text:
            blocks.append((title, text))
    join_txt = _format_sql_joins(sql)
    if join_txt:
        blocks.append(("JOIN", join_txt))
    where_blob = "\n".join(_where_lines(sql))
    if where_blob:
        blocks.append(("WHERE", where_blob))
    for title, key in (("GROUP BY", "group_by"), ("ORDER BY", "order_by")):
        text = _sql_clause(sql.get(key))
        if text:
            blocks.append((title, text))
    return blocks


def _has_meaningful_sel_sql(sql: Any) -> bool:
    """True when the structured ``sql`` object would render at least one clause."""
    return isinstance(sql, dict) and bool(_build_sel_sql_section_blocks(sql))


def _sel_sql_source(q: dict) -> tuple[str, Any]:
    """How to render a SELECT query: ("structured", sql) / ("raw", text) / ("none", None).

    Layer 2 usually leaves ``sql.select`` / ``sql.from`` empty and puts the whole
    statement in ``raw_sql``; then ``raw_sql`` is used.
    """
    sql = q.get("sql") if isinstance(q.get("sql"), dict) else {}
    raw = _normalize_sql_newlines(q.get("raw_sql") or "")
    if raw and not _sql_clause(sql.get("select")) and not _sql_clause(sql.get("from")):
        return "raw", raw
    if _has_meaningful_sel_sql(sql):
        return "structured", sql
    if raw:
        return "raw", raw
    return "none", None


def compose_sel_query_sql_text(q: dict) -> str:
    """Full SQL text of a query: ``raw_sql``, or the structured clauses one block each."""
    kind, payload = _sel_sql_source(q)
    if kind == "raw":
        return payload
    if kind == "structured":
        return "\n\n".join(
            f"{title}\n{_normalize_sql_newlines(body)}"
            for title, body in _build_sel_sql_section_blocks(payload)
        ).strip()
    return ""


# ─────────────────────────────────────────────
# SelQuery / UpdQuery sheets
# ─────────────────────────────────────────────


def _fill_sel_sql_layout(ws: Worksheet, sql: dict[str, Any], *,
                         last_column_block_row: int | None = None) -> int | None:
    """Paint SQL under the SELECT column grid as stacked 2-row blocks.

    Each clause (SELECT / FROM / JOIN / …) becomes a green header row plus a
    white body row (see ``apply_stacked_sql_blocks`` in this module). Placement starts at
    ``last_column_block_row + 1`` so it stays below expanded column rows.
    """
    if sql is None:
        return None

    last = last_column_block_row
    if last is None:
        r_select = find_row_by_cell_text(ws, "A", "SELECT句")
        if r_select is None:
            return None
        r_hdr = find_sel_query_column_header_row(ws, after_row=r_select, scan_rows=35)
        if r_hdr is None:
            return None
        tr = r_hdr + 1
        last = tr + template_span(ws, tr) - 1

    blocks = _build_sel_sql_section_blocks(sql)
    return apply_stacked_sql_blocks(
        ws,
        last + 1,
        blocks,
        col_start=_DD_STACKED_SQL_COL_START,
        col_end=_DD_STACKED_SQL_COL_END,
        body_col_start=_DD_STACKED_SQL_BODY_COL_START,
    )


def _fill_sel_raw_sql_layout(
    ws: Worksheet, raw_sql: str, *, last_column_block_row: int | None = None
) -> int | None:
    """Render raw SQL text under SELECT grid; one SQL line per worksheet row.

    Returns the last body row written, or ``None`` if nothing was written.
    """
    # Accept both real newlines and escaped literal sequences from upstream text.
    raw = _normalize_sql_newlines(raw_sql or "")
    if not raw:
        return None
    last = last_column_block_row
    if last is None:
        r_select = find_row_by_cell_text(ws, "A", "SELECT句")
        if r_select is None:
            return None
        r_hdr = find_sel_query_column_header_row(ws, after_row=r_select, scan_rows=35)
        if r_hdr is None:
            return None
        tr = r_hdr + 1
        last = tr + template_span(ws, tr) - 1
    lines = [ln for ln in raw.split("\n")]
    if not lines:
        lines = [raw]
    c1_letter = get_column_letter(_DD_STACKED_SQL_COL_START)
    c2_letter = get_column_letter(_DD_STACKED_SQL_COL_END)
    body_c1_letter = get_column_letter(_DD_STACKED_SQL_BODY_COL_START)
    hdr_fill = PatternFill(fill_type="solid", fgColor="C6EFCE")
    body_fill = PatternFill(fill_type="solid", fgColor="FFFFFF")
    hdr_font = Font(name=EXCEL_DOCUMENT_FONT_NAME, bold=True)
    body_align = Alignment(
        wrap_text=True, vertical="top", horizontal="left", indent=0,
    )
    hdr_align = Alignment(horizontal="left", vertical="center")
    thin_side = Side(style="thin", color="000000")
    row_border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side,
    )
    r = last + 1
    _unmerge_cells_touching(
        ws,
        r,
        r + len(lines),
        _DD_STACKED_SQL_COL_START,
        _DD_STACKED_SQL_COL_END,
    )
    ws.merge_cells(f"{c1_letter}{r}:{c2_letter}{r}")
    top = ws.cell(r, _DD_STACKED_SQL_COL_START)
    top.value = "RAW SQL"
    top.font = hdr_font
    top.fill = hdr_fill
    top.alignment = hdr_align
    top.border = row_border
    apply_excel_document_font(top)
    _apply_row_box_border(
        ws,
        r,
        col_start=_DD_STACKED_SQL_COL_START,
        col_end=_DD_STACKED_SQL_COL_END,
        side=thin_side,
    )
    for i, ln in enumerate(lines):
        r += 1
        ws.merge_cells(f"{body_c1_letter}{r}:{c2_letter}{r}")
        cell = ws.cell(r, _DD_STACKED_SQL_BODY_COL_START)
        write_text(cell, ln[:32000] if ln else None)
        cell.fill = body_fill
        cell.alignment = body_align
        cell.border = Border()
        _apply_row_outer_border(
            ws,
            r,
            col_start=_DD_STACKED_SQL_COL_START,
            col_end=_DD_STACKED_SQL_COL_END,
            side=thin_side,
            draw_top=(i == 0),
            draw_bottom=(i == len(lines) - 1),
        )
        apply_excel_document_font(cell)
    return r


def _fill_upd_sql_layout(
    ws: Worksheet,
    sql: dict[str, Any],
    *,
    upd_map: dict[str, str],
    stack_start_row: int,
) -> None:
    """Render UPD/INS SQL sections below the update grid as stacked SQL blocks."""
    if not sql:
        return

    clear_cols = tuple(
        sorted(
            set(upd_map.values()),
            key=lambda lt: column_index_from_string(lt),
        )
    )
    clear_block(ws, stack_start_row, stack_start_row + 40, clear_cols)

    blocks: list[tuple[str, str]] = []
    for title, key in (("SELECT", "select"), ("FROM", "from")):
        text = _sql_clause(sql.get(key))
        if text:
            blocks.append((title, text))
    jt = _format_sql_joins(sql)
    if jt:
        blocks.append(("JOIN", jt))
    where_lines = _where_lines(sql)
    if where_lines:
        blocks.append(("WHERE", "\n".join(where_lines)))
    group_by = _sql_clause(sql.get("group_by"))
    if group_by:
        blocks.append(("GROUP BY / HAVING", group_by))

    if not blocks:
        return
    # Same geometry as `_fill_sel_sql_layout` (A..AM merges), otherwise Upd stacks
    # from column B only and reads visually 'indented' vs sheet_4.
    _ = apply_stacked_sql_blocks(
        ws,
        stack_start_row,
        blocks,
        col_start=_DD_STACKED_SQL_COL_START,
        col_end=_DD_STACKED_SQL_COL_END,
        body_col_start=_DD_STACKED_SQL_BODY_COL_START,
    )


def _format_parameters_cell(params: Any) -> str | None:
    if params is None:
        return None
    if isinstance(params, list):
        s = "、".join(str(p) for p in params if str(p).strip())
        return s or None
    s = str(params).strip()
    return s or None


def _fill_query_common(ws: Worksheet, dd: dict, q: dict) -> None:
    """Fill the header block shared by SEL / UPD sheets (screen DD grid layout)."""
    fill_common_header(
        ws, dd,
        module_id=q.get("module_id"),
        module_name=q.get("module_name"),
    )
    put(
        ws,
        "E6",
        q.get("description") or q.get("overview") or q.get("query_purpose"),
    )
    put(ws, "E7", q.get("query_name"))
    put(ws, "E8", _format_parameters_cell(q.get("parameters")))


def _fill_report_compact_query_common(ws: Worksheet, dd: dict, q: dict) -> None:
    """Fill overview / Query名 / パラメータ rows on DD_Report compact SelQuery sheets."""
    fill_common_header(
        ws, dd,
        module_id=q.get("module_id"),
        module_name=q.get("module_name"),
    )
    desc = q.get("description") or q.get("overview") or q.get("query_purpose")
    r_gai = find_row_by_cell_text(ws, "A", "概要", row_start=3, row_end=12)
    if desc:
        put(ws, f"E{r_gai}", desc) if r_gai is not None else put(ws, "E5", desc)

    r_qn = find_row_by_cell_text(ws, "A", "Query名", row_start=3, row_end=15)
    if r_qn is not None:
        put(ws, f"E{r_qn}", q.get("query_name"))
    else:
        put(ws, "E6", q.get("query_name"))

    r_par = find_row_by_cell_text(ws, "A", "パラメータ", row_start=3, row_end=15)
    params_txt = _format_parameters_cell(q.get("parameters"))
    if r_par is not None:
        put(ws, f"E{r_par}", params_txt)
    else:
        put(ws, "E7", params_txt)


def _compact_report_sql_anchor_row(ws: Worksheet) -> int:
    """Row below ``パラメータ`` label (merged SQL band in DD_Report template)."""
    r = find_row_by_cell_text(ws, "A", "パラメータ", row_start=2, row_end=20)
    return (r + 1) if r is not None else 8


def _build_compact_report_sql_blocks(q: dict) -> list[tuple[str, str]]:
    """Build stacked SQL blocks for DD_Report compact SelQuery sheet."""
    kind, payload = _sel_sql_source(q)
    if kind == "structured":
        return _build_sel_sql_section_blocks(payload)
    if kind == "raw":
        return [("SQL", payload)]
    return []


def _fill_compact_report_sql_blocks(ws: Worksheet, q: dict) -> int | None:
    """Render SQL in compact report as green-key + white-body blocks."""
    r = _compact_report_sql_anchor_row(ws)
    col_start, col_end = 1, min(39, ws.max_column or 39)
    if col_start > col_end:
        return None
    blocks = _build_compact_report_sql_blocks(q)
    if not blocks:
        return None
    clear_block(ws, r, min(r + 120, ws.max_row), ("A",))
    return apply_stacked_sql_blocks(
        ws, r, blocks, col_start=col_start, col_end=col_end, body_col_start=2,
    )


def _is_dd_report_compact_sel_query_sheet(ws: Worksheet) -> bool:
    """DD_Report SelQuery sheets omit the ``SELECT句`` grid (single SQL band only)."""
    if find_row_by_cell_text(ws, "A", "SELECT句", row_start=1, row_end=50) is not None:
        return False
    return find_row_by_cell_text(ws, "A", "帳票ID", row_start=2, row_end=12) is not None


def _fill_sel_query_sql(ws: Worksheet, q: dict, last_col_r: int) -> int | None:
    """SQL below the SELECT grid: structured blocks, or ``raw_sql`` line by line."""
    kind, payload = _sel_sql_source(q)
    if kind == "structured":
        return _fill_sel_sql_layout(ws, payload, last_column_block_row=last_col_r)
    if kind == "raw":
        return _fill_sel_raw_sql_layout(ws, payload, last_column_block_row=last_col_r)
    return None


def fill_sel_query_sheet(ws: Worksheet, dd: dict, q: dict) -> None:
    """Fill a SelQuery sheet: DD_Report compact layout, or the screen ``SELECT句`` grid."""
    if _is_dd_report_compact_sel_query_sheet(ws):
        _fill_report_compact_query_common(ws, dd, q)
        r_sql = _compact_report_sql_anchor_row(ws)
        sql_last = _fill_compact_report_sql_blocks(ws, q)
        _stretch_sel_query_result_confirm_body(
            ws, body_style_row=r_sql - 1, sql_body_last_row=sql_last,
        )
        return

    _fill_query_common(ws, dd, q)

    # SELECT columns block: use the row directly under the grid header as the
    # style/data template; extra rows are inserted via expand_table.
    r_select = find_row_by_cell_text(ws, "A", "SELECT句")
    if r_select is None:
        return
    r_hdr = find_sel_query_column_header_row(ws, after_row=r_select, scan_rows=35)
    if r_hdr is None:
        return
    template_row = r_hdr + 1

    cols = q.get("columns", []) or q.get("select_columns", []) or []
    # stop purging at the next section header (FROM句 / WHERE句 etc.)
    purge_stale_data_rows(ws, template_row, SEL_COL.values(), max_scan=30)
    span0 = template_span(ws, template_row)
    if not cols:
        clear_row_cells(ws, template_row, SEL_COL.values())
        last_col_r = template_row + span0 - 1
    else:
        rows = expand_table(ws, template_row, len(cols))
        # Keep all generated SELECT data rows visually identical to the template row.
        # Some templates carry inconsistent borders on inserted middle rows.
        max_col = ws.max_column or 39
        for rr in rows:
            if rr != template_row:
                _duplicate_full_row_styles(ws, template_row, rr, max_col=max_col)
        for c, rr in zip(cols, rows):
            clear_row_cells(ws, rr, SEL_COL.values())
            if isinstance(c, dict):
                c = dict(c)
                # Standardized schema aliases.
                if "display_name" not in c and "alias" in c:
                    c["display_name"] = c["alias"]
                if "remark" not in c and "confirm_item" in c:
                    c["remark"] = c["confirm_item"]
            for key, col in SEL_COL.items():
                if key in c:
                    put(ws, f"{col}{rr}", c[key])

        for rr in rows:
            cv = ws.cell(rr, 3).value
            if cv is None:
                continue
            s = str(cv)
            if any(t in s for t in ("パラメータ", "Parameters")) or s.startswith(("AND A.", "A.")):
                clear_row_cells(ws, rr, ("C", "W", "AO"))
        last_col_r = rows[-1] + span0 - 1

    sql_last = _fill_sel_query_sql(ws, q, last_col_r)
    _stretch_sel_query_result_confirm_body(
        ws, body_style_row=last_col_r, sql_body_last_row=sql_last,
    )


def fill_upd_query_sheet(ws: Worksheet, dd: dict, q: dict) -> None:
    _fill_query_common(ws, dd, q)

    # UPDATE query has an extra row 9 with テーブル名 / テーブルID.
    sql_obj = q.get("sql") if isinstance(q.get("sql"), dict) else {}
    put(
        ws,
        "E9",
        q.get("table_physical_name")
        or q.get("table_name")
        or sql_obj.get("target_table")
        or sql_obj.get("update_table"),
    )
    put(
        ws,
        "Y9",
        q.get("table_id")
        or q.get("table_logical_id")
        or sql_obj.get("target_table"),
    )

    r_hdr = find_row_by_cell_text(ws, "A", "No", row_start=9, row_end=20)
    if r_hdr is None:
        return
    template_row = r_hdr + 1
    upd_map = _resolve_upd_col_map(ws, r_hdr)
    upd_cols = tuple(upd_map[k] for k in _UPD_FIELD_KEYS)

    cols = q.get("columns", []) or q.get("update_columns", []) or []
    if not cols and isinstance(q.get("sql"), dict):
        set_cols = q["sql"].get("set_columns", [])
        if isinstance(set_cols, list):
            cols = [
                {
                    "no": c.get("no") or str(i + 1),
                    "column_name": c.get("column"),
                    "source_area": c.get("source_table", ""),
                    "source_field": (
                        c.get("source_field_or_fixed_value")
                        or c.get("value_expression")
                    ),
                    "remark": c.get("notes") or c.get("confirm_item"),
                }
                for i, c in enumerate(set_cols)
                if isinstance(c, dict)
            ]
    purge_stale_data_rows(ws, template_row, upd_cols, max_scan=30)
    if not cols:
        clear_row_cells(ws, template_row, upd_cols)
        if isinstance(sql_obj, dict):
            _fill_upd_sql_layout(
                ws,
                sql_obj,
                upd_map=upd_map,
                stack_start_row=template_row + 1,
            )
        return
    rows = expand_table(ws, template_row, len(cols))
    # Keep all generated UPDATE data rows visually identical to the template row.
    max_col = ws.max_column or 39
    for rr in rows:
        if rr != template_row:
            _duplicate_full_row_styles(ws, template_row, rr, max_col=max_col)
    for c, rr in zip(cols, rows):
        clear_row_cells(ws, rr, upd_cols)
        for key in _UPD_FIELD_KEYS:
            col = upd_map[key]
            if key in c:
                put(ws, f"{col}{rr}", c[key])

    if isinstance(sql_obj, dict):
        # UpdQuery template row blocks span multiple visual rows; unlike SelQuery,
        # placement should follow the actual last populated data row.
        last_col_r = rows[-1]
        _fill_upd_sql_layout(
            ws,
            sql_obj,
            upd_map=upd_map,
            stack_start_row=last_col_r + 1,
        )


# ─────────────────────────────────────────────
# Sheet renaming / dispatchers / ordering
# ─────────────────────────────────────────────


def rename_sheets_to_screen_id(wb, tmpl_prefix: str, screen_id: str) -> None:
    """Replace ``tmpl_prefix`` with ``screen_id`` in all templated sheet names.

    Handles templated sheet name patterns:
        詳細仕様書_<prefix>_NN
        詳細仕様書_<prefix>      (report loose)
        <prefix>_NNSelQueryMM / <prefix>_SelQueryMM
        <prefix>_NNUpdQueryMM / <prefix>_UpdQueryMM
        <prefix>_NNInsQueryMM / <prefix>_InsQueryMM
    """
    if not tmpl_prefix or not screen_id or tmpl_prefix == screen_id:
        return
    for name in list(wb.sheetnames):
        new: str | None = None
        m = _RE_DETAIL.match(name)
        if m and _prefix_matches(tmpl_prefix, m.group("prefix")):
            new = _safe_title_preserve_suffix(
                "詳細仕様書_", screen_id, f"_{m.group('nn')}",
            )
        else:
            m = _RE_DETAIL_LOOSE.match(name)
            if (
                m
                and not _RE_DETAIL.match(name)
                and _prefix_matches(tmpl_prefix, m.group("id").strip())
            ):
                new = _safe_title_preserve_suffix("詳細仕様書_", screen_id, "")
            else:
                m = _RE_SELQUERY.match(name)
                if m and _prefix_matches(tmpl_prefix, m.group("prefix")):
                    new = _safe_title_preserve_suffix(
                        "", screen_id, f"_{m.group('nn')}SelQuery{m.group('mm')}",
                    )
                else:
                    m = _RE_SELQUERY_COMPACT.match(name)
                    if m and _prefix_matches(tmpl_prefix, m.group("prefix")):
                        new = _safe_title_preserve_suffix(
                            "", screen_id, f"_SelQuery{m.group('mm')}",
                        )
                    else:
                        m = _RE_UPDQUERY.match(name)
                        if m and _prefix_matches(tmpl_prefix, m.group("prefix")):
                            new = _safe_title_preserve_suffix(
                                "", screen_id, f"_{m.group('nn')}UpdQuery{m.group('mm')}",
                            )
                        else:
                            m = _RE_UPDQUERY_COMPACT.match(name)
                            if m and _prefix_matches(tmpl_prefix, m.group("prefix")):
                                new = _safe_title_preserve_suffix(
                                    "", screen_id, f"_UpdQuery{m.group('mm')}",
                                )
                            else:
                                m = _RE_INSQUERY.match(name)
                                if m and _prefix_matches(tmpl_prefix, m.group("prefix")):
                                    new = _safe_title_preserve_suffix(
                                        "", screen_id, f"_{m.group('nn')}InsQuery{m.group('mm')}",
                                    )
                                else:
                                    m = _RE_INSQUERY_COMPACT.match(name)
                                    if (
                                        m
                                        and _prefix_matches(tmpl_prefix, m.group("prefix"))
                                    ):
                                        new = _safe_title_preserve_suffix(
                                            "", screen_id, f"_InsQuery{m.group('mm')}",
                                        )
        if new is None:
            continue
        if new == name or new in wb.sheetnames:
            continue
        wb[name].title = new


def handle_detail_sheets(wb, dd: dict, screen_id: str,
                         *, purge_extras: bool = True,
                         detail_profile: str = "auto",
                         annotate_events: bool = False) -> None:
    """Fill 詳細仕様書_* sheets.

    ``detail_profile``:
        ``"auto"`` — legacy (``sheet_2_detail`` list clone if present).
        ``"form"`` — ``sheet_3_detail_specs`` / ``sheet_2_overview.details``;
            one **cloned** detail sheet per JSON item (same as module ``sheet_2_detail``).
            Ignores top-level ``sheet_2_detail`` (avoids ``[]`` wiping sheets).
        ``"module_macro"`` — prefer ``sheet_2_detail``; no overview.details
            fallback when ``sheet_3`` is absent; treats ``[]`` as absent.
    """
    source_key = "sheet_2_detail"
    raw: Any = None
    details: list[Any] = []

    if detail_profile == "form":
        source_key = "sheet_3_detail_specs"
        raw = dd.get("sheet_3_detail_specs")
        if raw is None and isinstance(dd.get("sheet_2_overview"), dict):
            source_key = "sheet_2_overview.details"
            raw = {"details": (dd.get("sheet_2_overview") or {}).get("details", [])}
        if raw is None:
            raw = {}
        raw = raw or {}
        if isinstance(raw, dict):
            details = raw.get("details", [])
        elif isinstance(raw, list):
            details = raw
        else:
            details = []
        if not isinstance(details, list):
            details = []
    else:
        raw = dd.get("sheet_2_detail")
        if raw == []:
            raw = None
        source_key = "sheet_2_detail"
        if raw is None:
            source_key = "sheet_3_detail_specs"
            raw = dd.get("sheet_3_detail_specs")
        if raw is None and detail_profile != "module_macro":
            source_key = "sheet_2_overview.details"
            raw = {"details": (dd.get("sheet_2_overview") or {}).get("details", [])}
        if raw is None:
            raw = {}
        raw = raw or {}
        if isinstance(raw, dict):
            details = raw.get("details", [])
        elif isinstance(raw, list):
            details = raw
        else:
            details = []
        if not isinstance(details, list):
            details = []

    sheets = _detail_sheets_in_order(wb, screen_id)

    # Form JSON: one 詳細仕様書 sheet per ``details[]`` row (clone first template).
    if (
        detail_profile == "form"
        and details
        and source_key in ("sheet_3_detail_specs", "sheet_2_overview.details")
    ):
        if not sheets:
            return
        tpl_name = sheets[0]
        used: set[str] = set()
        for idx, item in enumerate(details):
            if not isinstance(item, dict):
                continue
            target = _resolve_detail_sheet_name(item, screen_id, idx)
            if target in wb.sheetnames:
                ws = wb[target]
            else:
                ws = _copy_worksheet_safely(wb, wb[tpl_name])
                ws.title = target
            fill_detail_sheet(ws, dd, item, annotate_events=annotate_events)
            used.add(target)
        if purge_extras:
            for n in list(wb.sheetnames):
                m = _RE_DETAIL.match(n)
                if m and _prefix_matches(screen_id, m.group("prefix")) and n not in used:
                    del wb[n]
                    continue
                m = _RE_DETAIL_LOOSE.match(n)
                if (
                    m
                    and not _RE_DETAIL.match(n)
                    and _prefix_matches(screen_id, m.group("id").strip())
                    and n not in used
                ):
                    del wb[n]
        return

    # New module/macro schema: each item in sheet_2_detail corresponds to one
    # detail sheet; clone from the first detail template sheet (usually *_01).
    if (
        detail_profile != "form"
        and source_key == "sheet_2_detail"
        and isinstance(raw, list)
    ):
        if not sheets:
            return
        tpl_name = sheets[0]
        used: set[str] = set()
        for idx, item in enumerate(details):
            if not isinstance(item, dict):
                continue
            target = _resolve_detail_sheet_name(item, screen_id, idx)
            if target in wb.sheetnames:
                ws = wb[target]
            else:
                ws = _copy_worksheet_safely(wb, wb[tpl_name])
                ws.title = target
            fill_detail_sheet(ws, dd, item, annotate_events=annotate_events)
            used.add(target)
        if purge_extras:
            for n in list(wb.sheetnames):
                m = _RE_DETAIL.match(n)
                if m and _prefix_matches(screen_id, m.group("prefix")) and n not in used:
                    del wb[n]
                    continue
                m = _RE_DETAIL_LOOSE.match(n)
                if (
                    m
                    and not _RE_DETAIL.match(n)
                    and _prefix_matches(screen_id, m.group("id").strip())
                    and n not in used
                ):
                    del wb[n]
        return

    # Pair details[] with sheets[] in order (index-to-index).
    n = min(len(details), len(sheets))
    for i in range(n):
        fill_detail_sheet(wb[sheets[i]], dd, details[i] or {}, annotate_events=annotate_events)

    # Purge detail sheets with no matching JSON entry (optional).
    if purge_extras:
        for extra in sheets[n:]:
            del wb[extra]


def _resolve_query_sheet_name(q: dict, screen_id: str, *, kind: str,
                              fallback_idx: int, compact: bool = False) -> str:
    """Build the target sheet name for a query.

    Preference order:
        1. q["sheet_name"] if present (verbatim).
        2. q["sheet_name_pattern"] with {screen_id},{nn},{mm} substitutions.
        3. default: <screen_id>_00{Sel|Upd}Query{NN}   (NN from fallback_idx+1),
           or <screen_id>_{Sel|Upd}Query{NN} for compact (DD_Report) templates.
    """
    name = (q.get("sheet_name") or "").strip()
    if name:
        return _safe_title(name)

    pattern = (q.get("sheet_name_pattern") or "").strip()
    if pattern:
        nn = str(q.get("nn") or q.get("module_no") or "00").zfill(2)
        mm = str(q.get("mm") or q.get("seq") or (fallback_idx + 1)).zfill(2)
        name = (pattern
                .replace("{screen_id}", screen_id)
                .replace("{FormID}", screen_id)
                .replace("{nn}", nn)
                .replace("{mm}", mm))
        return _safe_title(name)

    suffix = "SelQuery" if kind == "sel" else "UpdQuery"
    nn = "" if compact else "00"
    return _safe_title_preserve_suffix(
        "",
        screen_id,
        f"_{nn}{suffix}{fallback_idx+1:02d}",
    )


def handle_query_sheets(wb, dd: dict, screen_id: str, *,
                         kind: str, purge_extras: bool = True) -> None:
    """Fill SelQuery or UpdQuery sheets according to the JSON input.

    When JSON provides no queries of a given kind, all template sheets of
    that kind are removed (``purge_extras=True``).
    """
    key  = "sheet_4_sel_query_templates" if kind == "sel" else "sheet_5_upd_query_templates"
    list_key = "sel_queries" if kind == "sel" else "upd_queries"
    fill_fn = fill_sel_query_sheet if kind == "sel" else fill_upd_query_sheet

    raw = dd.get(key) or {}
    if isinstance(raw, dict):
        queries = raw.get(list_key, []) or []
        singular = "sel_query" if kind == "sel" else "upd_query"
        if not queries and isinstance(raw.get(singular), dict):
            queries = [raw[singular]]
        if (
            not queries
            and kind == "upd"
            and isinstance(raw.get("ins_queries"), list)
        ):
            queries = raw["ins_queries"]
    elif isinstance(raw, list):
        # Standardized schema: sheet_4 / sheet_5 can be direct list.
        queries = raw
    else:
        queries = []
    if not isinstance(queries, list):
        queries = []

    template_sheets = [
        n for n in wb.sheetnames
        if (m := _match_query_sheet_name(n, kind=kind))
        and _prefix_matches(screen_id, m.group("prefix"))
    ]

    if not queries:
        if purge_extras:
            for n in template_sheets:
                del wb[n]
        return

    tpl_name = template_sheets[0] if template_sheets else None
    # DD_Report names query sheets <id>_SelQueryMM (no _NN module token).
    compact = tpl_name is not None and not (
        _RE_SELQUERY.match(tpl_name) or _RE_UPDQUERY.match(tpl_name) or _RE_INSQUERY.match(tpl_name)
    )
    used: set[str] = set()

    for idx, q in enumerate(queries):
        if not isinstance(q, dict):
            continue
        q = dict(q)
        if "description" not in q and "query_purpose" in q:
            q["description"] = q["query_purpose"]
        target = _resolve_query_sheet_name(q, screen_id, kind=kind,
                                            fallback_idx=idx, compact=compact)

        if target in wb.sheetnames:
            ws = wb[target]
        elif tpl_name is not None:
            ws = _copy_worksheet_safely(wb, wb[tpl_name])
            ws.title = target
        else:
            print(f"[WARN] No template {kind}-query sheet - cannot create: {target}")
            continue

        fill_fn(ws, dd, q)
        used.add(target)

    if purge_extras:
        for n in template_sheets:
            if n not in used:
                del wb[n]


def reorder_related_sheets(wb, screen_id: str) -> None:
    """Cluster detail/query sheets so each module's sheets are adjacent.

    Matches ``fill_dd_excel_frm.reorder_related_sheets`` plus DD_Report patterns:
    ``詳細仕様書_<id>`` without ``_NN`` and ``<id>_SelQueryMM`` (compact query names).
    """
    original = list(wb.sheetnames)

    detail_by_key: dict[str, str] = {}
    sel_by_key: dict[str, list[tuple[int, str]]] = {}
    upd_by_key: dict[str, list[tuple[int, str]]] = {}
    ins_by_key: dict[str, list[tuple[int, str]]] = {}
    related_set: set[str] = set()

    for name in original:
        m_mod = _RE_DETAIL_MODULE.match(name)
        if m_mod:
            key = m_mod.group("module")
            detail_by_key[key] = name
            related_set.add(name)
            continue
        m = _RE_DETAIL.match(name)
        if m and _prefix_matches(screen_id, m.group("prefix")):
            key = m.group("nn").zfill(2)
            detail_by_key[key] = name
            related_set.add(name)
            continue

        m_mod = _RE_QUERY_MODULE.match(name)
        if m_mod:
            key = m_mod.group("module")
            mm = int(m_mod.group("mm"))
            query_kind = m_mod.group("kind")
            if query_kind == "SelQuery":
                sel_by_key.setdefault(key, []).append((mm, name))
            elif query_kind == "UpdQuery":
                upd_by_key.setdefault(key, []).append((mm, name))
            else:
                ins_by_key.setdefault(key, []).append((mm, name))
            related_set.add(name)
            continue
        m = _RE_SELQUERY.match(name)
        if m and _prefix_matches(screen_id, m.group("prefix")):
            key = m.group("nn").zfill(2)
            mm = int(m.group("mm"))
            sel_by_key.setdefault(key, []).append((mm, name))
            related_set.add(name)
            continue

        m = _RE_UPDQUERY.match(name)
        if m and _prefix_matches(screen_id, m.group("prefix")):
            key = m.group("nn").zfill(2)
            mm = int(m.group("mm"))
            upd_by_key.setdefault(key, []).append((mm, name))
            related_set.add(name)
            continue

        m = _RE_INSQUERY.match(name)
        if m and _prefix_matches(screen_id, m.group("prefix")):
            key = m.group("nn").zfill(2)
            mm = int(m.group("mm"))
            ins_by_key.setdefault(key, []).append((mm, name))
            related_set.add(name)
            continue

        m_dl = _RE_DETAIL_LOOSE.match(name)
        if m_dl and not _RE_DETAIL.match(name):
            rid = m_dl.group("id").strip()
            if rid and _prefix_matches(screen_id, rid):
                detail_by_key[rid] = name
                related_set.add(name)
            continue

        m = _RE_SELQUERY_COMPACT.match(name)
        if m:
            pfx = m.group("prefix").strip()
            if pfx and _prefix_matches(screen_id, pfx):
                mm = int(m.group("mm"))
                sel_by_key.setdefault(pfx, []).append((mm, name))
                related_set.add(name)
            continue

        m = _RE_UPDQUERY_COMPACT.match(name)
        if m:
            pfx = m.group("prefix").strip()
            if pfx and _prefix_matches(screen_id, pfx):
                mm = int(m.group("mm"))
                upd_by_key.setdefault(pfx, []).append((mm, name))
                related_set.add(name)
            continue

        m = _RE_INSQUERY_COMPACT.match(name)
        if m:
            pfx = m.group("prefix").strip()
            if pfx and _prefix_matches(screen_id, pfx):
                mm = int(m.group("mm"))
                ins_by_key.setdefault(pfx, []).append((mm, name))
                related_set.add(name)

    if not related_set:
        return

    def _key_sort(k: str) -> tuple[int, str]:
        mm_a = re.match(r"^[A-Za-z](\d+)$", k)
        if mm_a:
            return (0, f"{k[0]}{int(mm_a.group(1)):06d}")
        if k.isdigit():
            return (1, f"{int(k):06d}")
        return (2, k)

    grouped: list[str] = []
    all_keys = sorted(
        set(detail_by_key) | set(sel_by_key) | set(upd_by_key) | set(ins_by_key),
        key=_key_sort,
    )
    for grp_key in all_keys:
        if grp_key in detail_by_key:
            grouped.append(detail_by_key[grp_key])
        grouped.extend(nm for _, nm in sorted(sel_by_key.get(grp_key, [])))
        grouped.extend(nm for _, nm in sorted(upd_by_key.get(grp_key, [])))
        grouped.extend(nm for _, nm in sorted(ins_by_key.get(grp_key, [])))

    anchor = min(i for i, n in enumerate(original) if n in related_set)
    ordered = (
        [n for n in original[:anchor] if n not in related_set]
        + grouped
        + [n for n in original[anchor:] if n not in related_set]
    )
    wb._sheets = [wb[n] for n in ordered]
