"""
fill_dd_excel_frm.py
--------------------
DD form JSON → Excel (``template/DD_Screen/dd_template.xlsx``) + CLI for
``inputs/DD_json/form``.

The fill engine (detail ``steps[]``, query sheets, stacked SQL blocks, …) is shared
with the report / macro / module tools in ``dd_common.py``; cell helpers come from
``fill_bd_excel``. ``detail_profile="form"`` (``sheet_2_overview``,
``sheet_3_detail_specs`` — one cloned detail sheet per item) is what the CLI and
``layer4/server.py`` use.
"""
from __future__ import annotations

import argparse
import json
import shutil
import sys
from pathlib import Path

_HERE = Path(__file__).resolve().parent
if str(_HERE) not in sys.path:
    sys.path.insert(0, str(_HERE))

try:
    from openpyxl import load_workbook
except ImportError as e:
    raise SystemExit("[ERR] Missing dependency. Run: pip install openpyxl") from e

from dd_common import (  # noqa: E402
    SHEET_OVERVIEW,
    detect_template_prefix,
    fill_overview,
    handle_detail_sheets,
    handle_query_sheets,
    has_meaningful_overview_payload,
    rename_sheets_to_screen_id,
    reorder_related_sheets,
)
from excel_translator import (  # noqa: E402
    restore_japanese_sheet_names,
    snapshot_template_labels,
    translate_workbook,
)
from fill_bd_excel import (  # noqa: E402
    SHEET_CHANGE_HISTORY,
    add_mock_change_history_args,
    build_mock_change_history,
    ensure_change_history_labels,
    fill_change_history,
)
from layer4_output_settings import get_dd_output_filename_prefix  # noqa: E402

_TOOLS_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_TEMPLATE = _TOOLS_ROOT / "template" / "DD_Screen" / "dd_template.xlsx"
DEFAULT_INPUT_DIR = _TOOLS_ROOT / "inputs" / "DD_json" / "form"
DEFAULT_OUTPUT_DIR = _TOOLS_ROOT / "outputs" / "DD_Output" / "form"


# ─────────────────────────────────────────────
# Main orchestrator
# ─────────────────────────────────────────────
def fill_excel(template_path: Path, dd: dict, output_path: Path, *,
               mock_change_history: bool = False,
               purge_extra_sheets: bool = True,
               detail_profile: str = "auto") -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(template_path, output_path)

    wb = load_workbook(output_path)
    restore_japanese_sheet_names(wb)
    if SHEET_CHANGE_HISTORY in wb.sheetnames:
        ensure_change_history_labels(wb[SHEET_CHANGE_HISTORY])
    labels = snapshot_template_labels(wb)

    tmpl_prefix = detect_template_prefix(wb.sheetnames)
    screen_id = (
        (dd.get("screen_id") or dd.get("form_id") or dd.get("name") or "").strip()
    )
    print(f"[INFO] Template prefix : {tmpl_prefix}")
    print(f"[INFO] JSON screen_id  : {screen_id}")

    # Rename sheets FIRST so downstream code can match by the new screen_id.
    if tmpl_prefix and screen_id:
        rename_sheets_to_screen_id(wb, tmpl_prefix, screen_id)
        active_prefix = screen_id
    else:
        active_prefix = tmpl_prefix or screen_id or ""

    # 1) 変更履歴 (no banner: rows 1-2 are the title and the table header)
    if SHEET_CHANGE_HISTORY in wb.sheetnames:
        ch_data = (dd.get("sheet_1_change_history") or {}).get("data") or []
        if not ch_data and mock_change_history:
            ch_data = build_mock_change_history(dd)
            print("[INFO] sheet_1_change_history.data is empty - using mock rows")
        fill_change_history(wb[SHEET_CHANGE_HISTORY], ch_data)
    else:
        print(f"[WARN] Sheet not found: {SHEET_CHANGE_HISTORY}")

    # 2) 概要仕様書
    if SHEET_OVERVIEW in wb.sheetnames:
        fill_overview(wb[SHEET_OVERVIEW], dd)
    elif has_meaningful_overview_payload(dd):
        print(f"[WARN] Sheet not found: {SHEET_OVERVIEW}")

    # 3) 詳細仕様書_* / 4) SelQuery / 5) UpdQuery sheets
    if active_prefix:
        handle_detail_sheets(
            wb,
            dd,
            active_prefix,
            purge_extras=purge_extra_sheets,
            detail_profile=detail_profile,
        )
        handle_query_sheets(wb, dd, active_prefix, kind="sel",
                            purge_extras=purge_extra_sheets)
        handle_query_sheets(wb, dd, active_prefix, kind="upd",
                            purge_extras=purge_extra_sheets)
        reorder_related_sheets(wb, active_prefix)

    translate_workbook(wb, labels)
    wb.save(output_path)


# --- CLI (form) ----------------------------------------------------------------


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Fill DD Excel from inputs/DD_json/form; write to outputs/DD_Output/form/.",
    )
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    parser.add_argument("--input-dir", type=Path, default=DEFAULT_INPUT_DIR)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--output", "-o", type=Path, default=None)
    parser.add_argument("--input", "-i", dest="inputs", nargs="*", default=None)
    add_mock_change_history_args(parser)
    parser.add_argument("--keep-extra-sheets", action="store_true")
    args = parser.parse_args()

    template_path = Path(args.template)
    if not template_path.is_file():
        raise SystemExit(f"[ERR] Template not found: {template_path}")

    if args.inputs:
        json_paths = [Path(p).resolve() for p in args.inputs]
        for jp in json_paths:
            if not jp.is_file():
                raise SystemExit(f"[ERR] Input JSON not found: {jp}")
    else:
        input_dir = Path(args.input_dir)
        if not input_dir.is_dir():
            raise SystemExit(
                f"[ERR] Input directory not found: {input_dir} (pass --input-dir or --input)"
            )
        json_paths = sorted(input_dir.glob("*.json"))

    if not json_paths:
        raise SystemExit(f"[ERR] No JSON (--input-dir {args.input_dir} *.json).")

    if args.output:
        if not args.inputs or len(args.inputs) != 1:
            raise SystemExit("[ERR] --output requires exactly one --input file.")
        jp = Path(args.inputs[0]).resolve()
        outp = Path(args.output)
        outp.parent.mkdir(parents=True, exist_ok=True)
        print(f"[INFO] Template : {template_path}")
        print(f"[INFO] Input    : {jp}")
        print(f"[INFO] Output   : {outp}")
        with open(jp, "r", encoding="utf-8") as f:
            dd = json.load(f)
        fill_excel(
            template_path,
            dd,
            outp,
            mock_change_history=args.mock_change_history,
            purge_extra_sheets=not args.keep_extra_sheets,
            detail_profile="form",
        )
        print(f"[DONE]  Saved -> {outp}")
        return

    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    for jp in json_paths:
        outp = output_dir / f"{get_dd_output_filename_prefix()}{jp.stem}.xlsx"
        print(f"[INFO] Template : {template_path}")
        print(f"[INFO] Input    : {jp}")
        print(f"[INFO] Output   : {outp}")
        with open(jp, "r", encoding="utf-8") as f:
            dd = json.load(f)
        fill_excel(
            template_path,
            dd,
            outp,
            mock_change_history=args.mock_change_history,
            purge_extra_sheets=not args.keep_extra_sheets,
            detail_profile="form",
        )
        print(f"[DONE]  Saved -> {outp}")


if __name__ == "__main__":
    main()
