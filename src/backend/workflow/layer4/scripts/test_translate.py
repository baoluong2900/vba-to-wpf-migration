import openpyxl
import re
from excel_translator import translate_workbook

from pathlib import Path
def test_translation():
    TEMPLATE_PATH = Path(__file__).resolve().parent / "../template/BD_Screen/bd_template.xlsx"
    wb = openpyxl.load_workbook(TEMPLATE_PATH)
    translate_workbook(wb)

    JP_REGEX = re.compile(r'[\u3040-\u309F\u30A0-\u30FF\u4E00-\u9FAF]')
    japanese_strings = []

    for name in wb.sheetnames:
        if JP_REGEX.search(name):
            japanese_strings.append(f"Sheet Name: {name}")
        ws = wb[name]
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and JP_REGEX.search(cell.value):
                    japanese_strings.append(f"{name}!{cell.coordinate}: {cell.value}")

    if japanese_strings:
        print("Found Japanese strings:")
        for s in japanese_strings[:20]:
            print(s)
        if len(japanese_strings) > 20:
            print(f"... and {len(japanese_strings)-20} more.")
    else:
        print("No Japanese strings found!")
