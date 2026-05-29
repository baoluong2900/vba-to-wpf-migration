import openpyxl
from pathlib import Path
import re
import sys

# Append scripts dir so we can import translator
sys.path.append('/Users/baoluong0209/GitTool/vba_mirgation/workflow/src/backend/workflow/layer4/scripts')
from excel_translator import translate_workbook

TEMPLATE_DIR = Path('/Users/baoluong0209/GitTool/vba_mirgation/workflow/src/backend/workflow/layer4/template')
JP_REGEX = re.compile(r'[\u3040-\u309F\u30A0-\u30FF\u4E00-\u9FAF]')

def clean_and_translate_template(file_path: Path):
    print(f"Processing: {file_path}")
    wb = openpyxl.load_workbook(file_path)
    
    # 1. Translate via translator
    translate_workbook(wb)
    
    # 2. Clear remaining dummy Japanese text and remove images
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        
        # Remove images
        if hasattr(ws, '_images'):
            ws._images = []
            
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and JP_REGEX.search(cell.value):
                    # For headers we might have missed, let's keep them if they are small,
                    # but clear the long dummy data.
                    val = cell.value
                    if "MONOS" in val or "出荷" in val or "補修" in val:
                        cell.value = ""
                    # If there's still some weird Japanese, we just wipe it to be safe 
                    # if it looks like dummy data. 
                    elif len(val) > 10 and not val.startswith("="):
                        cell.value = ""
                        
    wb.save(file_path)
    print(f"Saved: {file_path}")

for ext in ["**/*.xlsx"]:
    for f in TEMPLATE_DIR.rglob(ext):
        if not f.name.startswith("~"):
            clean_and_translate_template(f)

print("Done fixing templates!")
