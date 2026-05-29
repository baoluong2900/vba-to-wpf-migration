import openpyxl
import re
from pathlib import Path

def test_flogin():
    path = Path(__file__).resolve().parent / '../../../../excel_outputs/BD/screen/F_Login.xlsx'
    if not path.exists():
        print(f"File not found: {path}")
    else:
        wb = openpyxl.load_workbook(path)
        JP_REGEX = re.compile(r'[\u3040-\u309F\u30A0-\u30FF\u4E00-\u9FAF]')
        
        found = False
        for name in wb.sheetnames:
            if JP_REGEX.search(name):
                print(f"Sheet Name: {name}")
                found = True
            ws = wb[name]
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and JP_REGEX.search(cell.value):
                        print(f"{name}!{cell.coordinate}: {cell.value}")
                        found = True
        if not found:
            print("No Japanese text found in F_Login.xlsx!")
