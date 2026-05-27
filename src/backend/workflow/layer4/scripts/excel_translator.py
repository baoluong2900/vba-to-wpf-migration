import re
from openpyxl.workbook import Workbook

JP_REGEX = re.compile(r'[\u3040-\u309F\u30A0-\u30FF\u4E00-\u9FAF]')

JP_EN_MAPPING = {
    # Thay thế chung & Change History
    "変更履歴": "Change History",
    "改訂": "Revision",
    "初版": "First Edition",
    "更新内容": "Update Content",
    "更新日": "Update Date",
    "更新者": "Updated By",
    "作成者": "Author",
    "作成日": "Created Date",
    "備考": "Notes",
    "更新理由/案件番号": "Reason for Update / Project No",
    "システム名": "System Name",
    "サブシステム名": "Sub-system Name",
    "プロジェクト名": "Project Name",
    "承認者": "Approver",

    # Common terms
    "＜レコードソース＞": "Record Source",
    "レコードソース": "Record Source",
    "レコード": "Record",
    "処理概要": "Process Overview",
    "画面項目名": "Screen Item Name",
    "項目\n属性": "Field Attribute",
    "属性\n桁数": "Length",
    "出力仕様": "Output Specification",
    "●タイミング：": "●Timing:",
    "タイミング": "Timing",
    "内容": "Content",
    "処理": "Process",
    "書式": "Format",
    "（レイアウト）": "(Layout)",
    "(レイアウト)": "(Layout)",
    "（論理名）": "(Logical)",
    "（物理名）": "(Physical)",
    "（項目編集）": "(Field Edit)",
    "（画面編集）": "(Screen Edit)",
    "転送元": "Source",
    "ﾃｰﾌﾞﾙ名/領域名": "Table Name / Region",

    # Layout Sheet
    "画面仕様書(レイアウト)": "Screen Spec (Layout)",
    "画面仕様書 （レイアウト）": "Screen Spec (Layout)",
    "帳票仕様書(レイアウト)": "Report Spec (Layout)",
    "帳票仕様書 （レイアウト）": "Report Spec (Layout)",
    "画面仕様書(画面編集)": "Screen Spec (Edit)",
    "画面仕様書 （画面編集）": "Screen Spec (Edit)",
    "帳票仕様書（項目編集）": "Report Spec (Field Edit)",
    "画面仕様書（項目編集）": "Screen Spec (Field Edit)",
    "画面ID": "Screen ID",
    "画面名": "Screen Name",
    "帳票ID": "Report ID",
    "帳票名": "Report Name",
    "イメージ": "Image",
    "●機能概要：": "●Functional Overview:",
    "●制約事項：": "●Constraints:",
    "イベントID": "Event ID",
    "名前（論理）": "Logical Name",
    "名前（物理）": "Physical Name",
    "種類": "Control Type",
    "標題": "Caption/Label",
    "イベント": "Event Trigger",
    "遷移先": "Transition Target",

    # Edit/Field Sheet
    "画面項目名 （論理名）": "Screen Item (Logical Name)",
    "画面項目名 （物理名）": "Screen Item (Physical Name)",
    "項目 属性": "Field Attribute/Type",
    "属性 桁数": "Length",
    "必須": "Required",
    "桁数": "Digits",
    "形式": "Format",
    "入力仕様": "Input Specification",
    "その他入力値チェック": "Validation Rules/Checks",
    "条  件  及  び  編  集  方  法": "Conditions & Editing Methods",
    "ﾃｰﾌﾞﾙ名／領域名": "Table Name / Region",
    "テーブル名": "Table Name",
    "項目名": "Item Name",
    "転  送  元": "Source",
    "項目名／固定値": "Field Name / Fixed Value",
    "転送元　領域名": "Source Region Name",
    "転送元　項目名/固定値": "Source Field / Fixed Value",

    # Table Update Sheet (E***)
    "画面仕様書（テーブル更新）": "Screen Spec (Table Update)",
    "テーブル名(物理名)": "Table Name (Physical)",
    "テーブル(論理名)": "Table Name (Logical)",
    "テーブルID": "Table ID",
    "項目名（物理名）": "Field Name (Physical)",
    "項目名（論理名）": "Field Name (Logical)",

    # DD specific
    "プログラム・モジュール仕様書": "Program/Module Specification",
    "詳細仕様書": "Detailed Design Spec",
    "概要仕様書": "Overview Spec",
    "画面仕様書": "Screen Spec",
    "帳票仕様書": "Report Spec",
    "ファイル仕様書": "File Spec",
    "仕様書": "Spec",
    "モジュールID": "Module ID",
    "モジュール名": "Module Name",
    "概要": "Overview",
    "処理記述": "Process Description",
    "初期処理": "Initialization Process",
    "条件解除処理": "Condition Clear Process",
    "実績確定処理": "Result Commit Process",
    "入力データ欄": "Input Data Section",
    "表示名": "Display Name",
    "カラム名": "Column Name",
    "Query名": "Query Name",
    "SELECT句": "SELECT Clause",
    "パラメータ": "Parameters",
    
    # Detailed mappings to fix mixed Japanese-English replacements
    "パラメータ情報": "Parameters Info",
    "引数情報": "Arguments Info",
    "項目名称": "Item Name / Label",
    "シンボル名": "Symbol Name",
    "戻り値": "Return Value",
    "全体桁": "Total Digits",
    "小数桁": "Decimal Digits",
    "特記事項": "Special Notes",
    "制約事項": "Constraints",
    "機能概要": "Functional Overview",
    "名称": "Name",
    "引数": "Arguments",
    "説明": "Description",
    "型": "Type",
    "コード": "Code",
    "タイプ": "Type",
    "サブ": "Sub",
    "情報": "Information",
}

def translate_workbook(wb: Workbook):
    try:
        from config import TARGET_LANGUAGE
    except ImportError:
        # Fallback if imported from outer python script
        TARGET_LANGUAGE = "en"
        
    if TARGET_LANGUAGE.lower() not in ("en", "english"):
        return
        
    # Sort mapping by key length in descending order to avoid partial replacement issues
    sorted_mapping = sorted(JP_EN_MAPPING.items(), key=lambda x: len(x[0]), reverse=True)
        
    # Translate sheet names
    for name in list(wb.sheetnames):
        ws = wb[name]
        new_name = name
        for jp, en in sorted_mapping:
            if jp in new_name:
                new_name = new_name.replace(jp, en)
        # Limit to 31 chars (Excel sheet name constraint) and clean special chars
        new_name = re.sub(r'[\\*?:/\[\]]', '', new_name)[:31].strip()
        if new_name != name and new_name not in wb.sheetnames:
            try:
                ws.title = new_name
            except Exception:
                pass
                
    # Translate cell values
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        # Iterate over all non-empty cells
        for row in ws.iter_rows():
            for cell in row:
                val = cell.value
                if isinstance(val, str) and JP_REGEX.search(val):
                    val_clean = val
                    for jp, en in sorted_mapping:
                        if jp in val_clean:
                            val_clean = val_clean.replace(jp, en)
                    cell.value = val_clean
