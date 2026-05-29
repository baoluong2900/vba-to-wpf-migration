"""
fill_bd_report_excel.py
-----------------------
Fill BD JSON for **帳票 (report)** into ``bd_report_template.xlsx``.

Template:
    02_tools/template/BD_Report/bd_report_template.xlsx

Sheets:
    - 変更履歴
    - 帳票仕様書(レイアウト)
    - 帳票仕様書（項目編集）

JSON shape matches screen BD (``sheet_1_change_history``, ``sheet_2_layout``,
``sheet_3_screen_edit`` optional, ``sheet_4_table_update``, ``sheet_5_file_spec``).

``sheet_5_file_spec`` rows are written into **帳票仕様書（項目編集）** (same table as
``sheet_3_screen_edit``), not a separate ``ファイル仕様書`` sheet — output stays **3 worksheets**.

Defaults (under 02_tools):
    template   : template/BD_Report/bd_report_template.xlsx
    input dir  : inputs/BD_json/report
    output dir : outputs/BD_Output/report

Layout sheet image (sheet 2 ``帳票仕様書(レイアウト)``): resolve under
``Access_DesignView/Report`` first, then ``Access_DesignView/Form``.
Use ``--image-dir`` to prepend extra folders.

Usage::

    cd 02_tools
    python -X utf8 scripts/fill_bd_report_excel.py

    python -X utf8 scripts/fill_bd_report_excel.py \\
        --input  inputs/BD_json/report/R005_作業報告書_弊社控え.json \\
        --output outputs/BD_Output/report/R005_作業報告書_弊社控え.xlsx
"""
from __future__ import annotations

import argparse
import json
import shutil
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
    from openpyxl.styles import Border, Side
    from openpyxl.worksheet.worksheet import Worksheet
except ImportError as e:
    raise SystemExit("[ERR] Missing dependency. Run: pip install openpyxl") from e

_HERE = Path(__file__).resolve().parent
if str(_HERE) not in sys.path:
    sys.path.insert(0, str(_HERE))

from layer4_output_settings import get_bd_report_output_filename_prefix  # noqa: E402
from fill_bd_excel import (  # noqa: E402
    DEFAULT_IMAGE_DIR,
    build_mock_change_history,
    clear_row_cells,
    expand_table,
    fill_branded_header_fields_by_label,
    fill_change_history,
    find_row_by_cell_contains,
    find_row_by_cell_text,
    handle_event_sheets,
    locate_template_row,
    purge_stale_data_rows,
    put,
    replace_layout_screen_image,
    _resolve_screen_image,
)

_TOOLS_ROOT = Path(__file__).resolve().parent.parent

DEFAULT_REPORT_TEMPLATE = _TOOLS_ROOT / "template" / "BD_Report" / "bd_report_template.xlsx"
DEFAULT_INPUT_DIR = _TOOLS_ROOT / "inputs" / "BD_json" / "report"
DEFAULT_OUTPUT_DIR = _TOOLS_ROOT / "outputs" / "BD_Output" / "report"

# Report layout images (prioritise Report over Form — not the same as screen BD)
DEFAULT_REPORT_IMAGE_DIR = _TOOLS_ROOT / "Access_DesignView" / "Report"


def _report_image_search_dirs(extra: list[Path] | None) -> list[Path]:
    """Search order: ``extra`` (CLI --image-dir), Report, Form."""
    base: list[Path] = [DEFAULT_REPORT_IMAGE_DIR, DEFAULT_IMAGE_DIR]
    if not extra:
        return base
    seen: set[str] = set()
    out: list[Path] = []
    for p in extra + base:
        key = str(p.resolve()) if p.exists() else str(p)
        if key not in seen:
            seen.add(key)
            out.append(p)
    return out

SHEET_CHANGE_HISTORY = "変更履歴"
SHEET_REPORT_LAYOUT = "帳票仕様書(レイアウト)"
SHEET_REPORT_EDIT = "帳票仕様書（項目編集）"

# 帳票仕様書（項目編集） data row: № / 項目名 / 転送元(表) / 転送元(項目) / …
# (differs from 画面仕様書(画面編集) which uses AH/AN for source columns.)
REPORT_EDIT_FALLBACK = 9
REPORT_EDIT_COL = {
    "item_no":                     "A",
    "logical_name":                "C",
    "source_table":                "K",
    "source_field_or_fixed_value": "S",
    "condition_and_format":        "AA",
    "chk":                         "AR",
    "notes":                       "AT",
}
REPORT_EDIT_NOTES_RIGHT_EDGE_COL = "BB"

LAYOUT_RECSRC_COL = "AN"
LAYOUT_RECSRC_MARKERS = [("AM", "＜レコードソース＞")]
LAYOUT_RECSRC_FALLBACK = 11

LAYOUT_PROC_MARKERS = [("AM", "【処理概要】")]
LAYOUT_PROC_FALLBACK = 23

from fill_bd_excel import LAYOUT_EVENT_COL, LAYOUT_PROC_COL  # noqa: E402


def fill_common_header_report_layout(ws: Worksheet, bd: dict) -> None:
    """Header on 帳票仕様書(レイアウト)."""
    fill_branded_header_fields_by_label(ws, author=bd.get("author") or bd.get("updated_by"))
    put(ws, "G2", bd.get("subsystem_name"))
    put(ws, "G4", bd.get("screen_id"))
    put(ws, "S4", bd.get("screen_name"))
    put(ws, "AK2", bd.get("updated_by"))
    put(ws, "AR2", bd.get("updated_date"))


def fill_common_header_report_edit(ws: Worksheet, bd: dict) -> None:
    """Header on 帳票仕様書（項目編集）."""
    fill_branded_header_fields_by_label(ws, author=bd.get("author") or bd.get("updated_by"))
    put(ws, "G2", bd.get("subsystem_name"))
    put(ws, "G4", bd.get("screen_id"))
    put(ws, "S4", bd.get("screen_name"))
    put(ws, "AN2", bd.get("updated_by"))
    put(ws, "AW2", bd.get("updated_date"))


def locate_report_edit_template_row(ws: Worksheet, fallback: int = 9) -> int:
    """Row below 転送元 sub-header ``K = ﾃｰﾌﾞﾙ名／領域名``; else ＜詳細＞; else screen edit locator."""
    r = find_row_by_cell_text(ws, "K", "ﾃｰﾌﾞﾙ名／領域名")
    if r is not None:
        return r + 1
    r = find_row_by_cell_text(ws, "C", "＜詳細＞")
    if r is not None:
        return r + 1
    from fill_bd_excel import locate_edit_template_row
    return locate_edit_template_row(ws, fallback=fallback)


def _normalize_report_edit_item(raw: dict) -> dict:
    """Map ``sheet_3`` / ``sheet_5`` keys onto ``REPORT_EDIT_COL``."""
    return {
        "item_no": raw.get("item_no") or raw.get("no"),
        "logical_name": (
            raw.get("logical_name")
            or raw.get("output_field_name")
            or raw.get("physical_name")
            or raw.get("target_field_physical")
        ),
        "source_table": raw.get("source_table") or raw.get("source_area"),
        "source_field_or_fixed_value": (
            raw.get("source_field_or_fixed_value") or raw.get("source_field")
        ),
        "condition_and_format": (
            raw.get("condition_and_format")
            or raw.get("condition")
            or raw.get("other_input_checks")
        ),
        "chk": raw.get("chk"),
        "notes": raw.get("notes") or raw.get("remark"),
    }


def collect_report_edit_items(bd: dict) -> list[dict]:
    """Merge ``sheet_3_screen_edit.items`` and all ``sheet_5_file_spec`` blocks."""
    out: list[dict] = []
    edit = bd.get("sheet_3_screen_edit", {}) or {}
    for it in edit.get("items", []) or []:
        if isinstance(it, dict):
            out.append(_normalize_report_edit_item(it))

    raw5 = bd.get("sheet_5_file_spec")
    blocks: list[dict] = []
    if isinstance(raw5, list):
        blocks = [b for b in raw5 if isinstance(b, dict)]
    elif isinstance(raw5, dict):
        if raw5.get("header") is not None or raw5.get("items") is not None:
            blocks = [raw5]
    for block in blocks:
        for it in block.get("items", []) or []:
            if isinstance(it, dict):
                out.append(_normalize_report_edit_item(it))
    return out


def fill_report_edit(ws: Worksheet, bd: dict) -> None:
    """Fill 帳票仕様書（項目編集） from ``sheet_3`` + ``sheet_5`` (one physical table)."""
    thin_right = Side(style="thin", color="000000")

    def _force_notes_right_border(row_start: int, row_end: int) -> None:
        from openpyxl.utils import column_index_from_string

        notes_col = column_index_from_string(REPORT_EDIT_COL["notes"])
        edge_col = column_index_from_string(REPORT_EDIT_NOTES_RIGHT_EDGE_COL)
        for rr in range(row_start, row_end + 1):
            for cc in (notes_col, edge_col):
                cell = ws.cell(rr, cc)
                b = cell.border
                cell.border = Border(
                    left=b.left,
                    right=thin_right,
                    top=b.top,
                    bottom=b.bottom,
                    diagonal=b.diagonal,
                    diagonal_direction=b.diagonal_direction,
                    outline=b.outline,
                    vertical=b.vertical,
                    horizontal=b.horizontal,
                )
        # For merged 備考 header/data ranges, the effective right border can be
        # inherited from the top-left cell style. Normalize those as well.
        for rg in ws.merged_cells.ranges:
            if rg.max_col != edge_col:
                continue
            if rg.max_row < row_start or rg.min_row > row_end:
                continue
            tl = ws.cell(rg.min_row, rg.min_col)
            b = tl.border
            tl.border = Border(
                left=b.left,
                right=thin_right,
                top=b.top,
                bottom=b.bottom,
                diagonal=b.diagonal,
                diagonal_direction=b.diagonal_direction,
                outline=b.outline,
                vertical=b.vertical,
                horizontal=b.horizontal,
            )

    fill_common_header_report_edit(ws, bd)
    items = collect_report_edit_items(bd)
    template_row = locate_report_edit_template_row(ws, fallback=REPORT_EDIT_FALLBACK)
    cols = list(REPORT_EDIT_COL.values())
    notes_hdr_row = find_row_by_cell_contains(
        ws, REPORT_EDIT_COL["notes"], "備",
        row_start=max(1, template_row - 5),
        row_end=template_row + 2,
    )
    if notes_hdr_row is None:
        notes_hdr_row = max(1, template_row - 2)
    if not items:
        clear_row_cells(ws, template_row, cols)
        _force_notes_right_border(notes_hdr_row, template_row)
        return
    rows = expand_table(ws, template_row, len(items))
    for row_data, r in zip(items, rows):
        clear_row_cells(ws, r, cols)
        for key, col in REPORT_EDIT_COL.items():
            val = row_data.get(key)
            if val is not None and str(val).strip() != "":
                put(ws, f"{col}{r}", val)
    _force_notes_right_border(notes_hdr_row, rows[-1])


def fill_report_layout(ws: Worksheet, bd: dict,
                       image_search_dirs: list[Path] | None = None) -> None:
    layout = bd.get("sheet_2_layout", {}) or {}
    fill_common_header_report_layout(ws, bd)

    put(ws, "E6", layout.get("description") or bd.get("overview"))
    put(ws, "AK6", layout.get("timing"))
    put(ws, "E7", layout.get("constraints"))

    search_dirs = (
        list(image_search_dirs)
        if image_search_dirs
        else _report_image_search_dirs(None)
    )
    resolved = _resolve_screen_image(
        layout.get("screen_image"),
        bd.get("screen_id"),
        search_dirs,
    )
    if resolved is not None:
        replace_layout_screen_image(ws, resolved)
        print(f"[INFO] Replaced report layout image with: {resolved}")
    else:
        img = layout.get("screen_image")
        if img and isinstance(img, str):
            put(ws, "A9", img)

    recsrc_row = locate_template_row(
        ws, LAYOUT_RECSRC_MARKERS, offset=1, fallback=LAYOUT_RECSRC_FALLBACK,
    )
    rec_sources = layout.get("record_sources", []) or []
    purge_stale_data_rows(ws, recsrc_row, [LAYOUT_RECSRC_COL])
    if not rec_sources:
        clear_row_cells(ws, recsrc_row, [LAYOUT_RECSRC_COL])
    else:
        rows = expand_table(ws, recsrc_row, len(rec_sources))
        for src, r in zip(rec_sources, rows):
            clear_row_cells(ws, r, [LAYOUT_RECSRC_COL])
            val = src.get("query_display") if isinstance(src, dict) else src
            put(ws, f"{LAYOUT_RECSRC_COL}{r}", val)

    r_proc_hdr = find_row_by_cell_text(ws, "AM", "【処理概要】")
    if r_proc_hdr is not None:
        proc_row = r_proc_hdr + 1
        proc_list = layout.get("process_overview", []) or []
        purge_stale_data_rows(ws, proc_row, [LAYOUT_PROC_COL], max_scan=25)
        if not proc_list:
            clear_row_cells(ws, proc_row, [LAYOUT_PROC_COL])
        else:
            rows = expand_table(ws, proc_row, len(proc_list))
            for proc, r in zip(proc_list, rows):
                clear_row_cells(ws, r, [LAYOUT_PROC_COL])
                val = proc.get("summary_line") if isinstance(proc, dict) else proc
                put(ws, f"{LAYOUT_PROC_COL}{r}", val)

    grp = (layout.get("grouping") or "").strip()
    srt = (layout.get("sorting") or "").strip()
    if grp or srt:
        r_g = find_row_by_cell_contains(ws, "AM", "＜グループ化・並べ替え＞")
        if r_g is not None:
            summary = " / ".join(x for x in (grp, srt) if x)
            put(ws, f"AO{r_g + 2}", summary)

    r_ev_hdr = find_row_by_cell_text(ws, "A", "イベントID")
    if r_ev_hdr is not None:
        events_row = r_ev_hdr + 1
        events = layout.get("events", []) or []
        if not events:
            clear_row_cells(ws, events_row, LAYOUT_EVENT_COL.values())
        else:
            rows = expand_table(ws, events_row, len(events))
            for ev, r in zip(events, rows):
                clear_row_cells(ws, r, LAYOUT_EVENT_COL.values())
                ev_copy = dict(ev)
                ev_copy.pop("action_type", None)
                for key, col in LAYOUT_EVENT_COL.items():
                    if key in ev_copy:
                        put(ws, f"{col}{r}", ev_copy[key])


def fill_report_excel(
    template_path: Path,
    bd: dict,
    output_path: Path,
    *,
    mock_change_history: bool = True,
    image_search_dirs: list[Path] | None = None,
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(template_path, output_path)

    wb = load_workbook(output_path)

    try:
        import sys
        from pathlib import Path
        _scripts_dir = Path(__file__).resolve().parent
        if str(_scripts_dir) not in sys.path:
            sys.path.insert(0, str(_scripts_dir))
        from excel_translator import restore_japanese_sheet_names
        restore_japanese_sheet_names(wb)
    except Exception as e:
        print(f"[WARN] Failed to restore Japanese sheet names: {e}")

    if SHEET_CHANGE_HISTORY in wb.sheetnames:
        ch_data = (bd.get("sheet_1_change_history", {}) or {}).get("data", []) or []
        if not ch_data and mock_change_history:
            ch_data = build_mock_change_history(bd)
            print("[INFO] sheet_1_change_history.data is empty - using mock rows")
        fill_change_history(wb[SHEET_CHANGE_HISTORY], ch_data)
    else:
        print(f"[WARN] Sheet not found: {SHEET_CHANGE_HISTORY}")

    img_dirs = _report_image_search_dirs(
        list(image_search_dirs) if image_search_dirs else None,
    )
    if SHEET_REPORT_LAYOUT in wb.sheetnames:
        fill_report_layout(wb[SHEET_REPORT_LAYOUT], bd, image_search_dirs=img_dirs)
    else:
        print(f"[WARN] Sheet not found: {SHEET_REPORT_LAYOUT}")

    if SHEET_REPORT_EDIT in wb.sheetnames:
        fill_report_edit(wb[SHEET_REPORT_EDIT], bd)
    else:
        print(f"[WARN] Sheet not found: {SHEET_REPORT_EDIT}")

    # sheet_5_file_spec is merged into 帳票仕様書（項目編集） — do not add ファイル仕様書.
    handle_event_sheets(wb, bd)

    try:
        import sys
        from pathlib import Path
        _scripts_dir = Path(__file__).resolve().parent
        if str(_scripts_dir) not in sys.path:
            sys.path.insert(0, str(_scripts_dir))
        from excel_translator import translate_workbook
        translate_workbook(wb)
    except Exception as e:
        print(f"[WARN] Failed to translate workbook: {e}")

    wb.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Fill BD report JSON into bd_report_template.xlsx.",
    )
    parser.add_argument(
        "--template",
        type=Path,
        default=DEFAULT_REPORT_TEMPLATE,
        help=f"Report Excel template (default: {DEFAULT_REPORT_TEMPLATE})",
    )
    parser.add_argument("--input", default=None, help="Single JSON (with --output).")
    parser.add_argument("--output", default=None, help="Single .xlsx output.")
    parser.add_argument(
        "--input-dir",
        type=Path,
        default=None,
        help=f"Batch input folder (default if no --input: {DEFAULT_INPUT_DIR})",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=None,
        help=f"Batch output folder (default: {DEFAULT_OUTPUT_DIR})",
    )
    parser.add_argument(
        "--no-mock-change-history",
        dest="mock_change_history",
        action="store_false",
        default=True,
        help="Disable auto-fill of 変更履歴 mock rows when JSON data is empty.",
    )
    parser.add_argument(
        "--image-dir",
        action="append",
        default=None,
        help="Prepend these folders when resolving layout images "
             f"(after that: {DEFAULT_REPORT_IMAGE_DIR}, then Form).",
    )
    args = parser.parse_args()

    template_path = Path(args.template)
    if not template_path.is_file():
        raise SystemExit(f"[ERR] Template not found: {template_path}")

    extra_img = [Path(p) for p in args.image_dir] if args.image_dir else None
    image_dirs = _report_image_search_dirs(extra_img)

    if args.input is None:
        in_dir = args.input_dir or DEFAULT_INPUT_DIR
        out_dir = args.output_dir or DEFAULT_OUTPUT_DIR
        if not in_dir.is_dir():
            raise SystemExit(f"[ERR] Input directory not found: {in_dir}")
        json_files = sorted(in_dir.glob("*.json"))
        if not json_files:
            raise SystemExit(f"[ERR] No *.json in {in_dir}")
        out_dir.mkdir(parents=True, exist_ok=True)
        print(f"[INFO] Template : {template_path}")
        print(f"[INFO] Input dir : {in_dir} ({len(json_files)} file(s))")
        print(f"[INFO] Output dir: {out_dir}")
        ok = 0
        failed: list[tuple[str, str]] = []
        for jp in json_files:
            op = out_dir / f"{get_bd_report_output_filename_prefix()}{jp.stem}.xlsx"
            try:
                with open(jp, "r", encoding="utf-8") as f:
                    bd = json.load(f)
                fill_report_excel(
                    template_path,
                    bd,
                    op,
                    mock_change_history=args.mock_change_history,
                    image_search_dirs=image_dirs,
                )
                print(f"[DONE]  {jp.name} -> {op.name}")
                ok += 1
            except json.JSONDecodeError as e:
                failed.append((jp.name, str(e)))
                print(f"[ERR]   {jp.name}: invalid JSON — {e}")
            except Exception as e:
                failed.append((jp.name, str(e)))
                print(f"[ERR]   {jp.name}: {e}")
        print(f"[SUMMARY] {ok} ok, {len(failed)} failed")
        if failed:
            raise SystemExit(1)
        return

    if not args.output:
        raise SystemExit("[ERR] --output required with --input.")
    ip, op = Path(args.input), Path(args.output)
    if not ip.is_file():
        raise SystemExit(f"[ERR] Input not found: {ip}")
    with open(ip, "r", encoding="utf-8") as f:
        bd = json.load(f)
    print(f"[INFO] Template: {template_path}")
    print(f"[INFO] Input   : {ip}")
    print(f"[INFO] Output  : {op}")
    fill_report_excel(
        template_path,
        bd,
        op,
        mock_change_history=args.mock_change_history,
        image_search_dirs=image_dirs,
    )
    print(f"[DONE]  Saved -> {op}")


if __name__ == "__main__":
    main()
