"""
llm/cost_tracker.py — Track token usage and estimated cost.

Singleton: get_tracker() returns a shared instance within a session.
"""
from __future__ import annotations

import json
import logging
import os
import threading
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from config import LOG_DIR

logger = logging.getLogger(__name__)

# Serialises the read-modify-write of token_log.json / token_log.xlsx (one gateway process, many threads)
_REPORT_LOCK = threading.Lock()

# Estimated cost in USD / 1M tokens (separate input / output)
_PRICE_INPUT_PER_1M: dict[str, float] = {
    "gemini-2.0-flash":     0.10,
    "gemini-1.5-flash":     0.10,
    "gemini-2.5-pro":       1.25,
    "gpt-4o":               2.50,
    "gpt-4o-mini":          0.15,
    "claude-3-5-sonnet":    3.00,
    "deepseek-v3":          0.27,
    "deepseek-chat":        0.14,
    "deepseek-reasoner":    0.55,
    "qwen-max":             2.80,
    "qwen-max-latest":      2.80,  # approx (= qwen-max)
    "qwen-plus":            0.56,
    "qwen-plus-2025-07-28": 0.56,  # approx (= qwen-plus)
    "qwen-turbo":           0.04,
    "qwen-long":            0.07,
    "qwen3-coder-instruct": 0.14,
    "qwen3-coder-480b-a35b-instruct": 2.80,  # approx (= qwen-max)
    "qwen3-coder-next":     0.14,  # approx (= qwen3-coder-instruct)
    "qwen3-235b-a22b-instruct-2507": 0.56,  # approx (= qwen-plus)
    "default":              1.00,
}

_PRICE_OUTPUT_PER_1M: dict[str, float] = {
    "gemini-2.0-flash":     0.40,
    "gemini-1.5-flash":     0.40,
    "gemini-2.5-pro":       5.00,
    "gpt-4o":               10.0,
    "gpt-4o-mini":          0.60,
    "claude-3-5-sonnet":    15.0,
    "deepseek-v3":          1.10,
    "deepseek-chat":        0.28,
    "deepseek-reasoner":    2.19,
    "qwen-max":             8.40,
    "qwen-max-latest":      8.40,  # approx (= qwen-max)
    "qwen-plus":            1.68,
    "qwen-plus-2025-07-28": 1.68,  # approx (= qwen-plus)
    "qwen-turbo":           0.11,
    "qwen-long":            0.21,
    "qwen3-coder-instruct": 0.56,
    "qwen3-coder-480b-a35b-instruct": 8.40,  # approx (= qwen-max)
    "qwen3-coder-next":     0.56,  # approx (= qwen3-coder-instruct)
    "qwen3-235b-a22b-instruct-2507": 1.68,  # approx (= qwen-plus)
    "default":              2.00,
}


def _move_aside(path: Path, exc: Exception) -> bool:
    """Rename an unreadable report file to <name>.<timestamp>.bak so its history is never overwritten."""
    backup = path.with_name(f"{path.name}.{datetime.now():%Y%m%d_%H%M%S_%f}.bak")
    try:
        os.replace(path, backup)
    except OSError as move_exc:
        logger.error("[CostTracker] %s is unreadable (%s) and could not be moved aside (%s) — report not saved",
                     path.name, exc, move_exc)
        return False
    logger.warning("[CostTracker] %s was unreadable (%s); kept as %s and started a new file",
                   path.name, exc, backup.name)
    return True


@dataclass
class TokenRecord:
    item_name: str
    model_id: str
    stage: str
    input_tokens: int
    output_tokens: int
    cost_usd: float
    flow_type: str = "Unknown"


@dataclass
class CostTracker:
    """Tracks token usage and estimated cost for a pipeline run."""
    records: list[TokenRecord] = field(default_factory=list)
    last_model_id: str = "Unknown"

    @property
    def total_input_tokens(self) -> int:
        return sum(r.input_tokens for r in self.records)

    @property
    def total_output_tokens(self) -> int:
        return sum(r.output_tokens for r in self.records)

    @property
    def total_tokens(self) -> int:
        return self.total_input_tokens + self.total_output_tokens

    @property
    def total_cost(self) -> float:
        return sum(r.cost_usd for r in self.records)

    def record(
        self,
        item_name: str,
        model_id: str,
        stage: str,
        input_tokens: int,
        output_tokens: int,
        flow_type: str = "Unknown",
    ) -> None:
        rate_in  = _PRICE_INPUT_PER_1M.get(model_id, _PRICE_INPUT_PER_1M["default"])
        rate_out = _PRICE_OUTPUT_PER_1M.get(model_id, _PRICE_OUTPUT_PER_1M["default"])
        
        self.last_model_id = model_id
        cost = (max(1, input_tokens) / 1_000_000 * rate_in) + (max(1, output_tokens) / 1_000_000 * rate_out)
        rec = TokenRecord(
            item_name=item_name,
            model_id=model_id,
            stage=stage,
            input_tokens=input_tokens,
            output_tokens=output_tokens,
            cost_usd=cost,
            flow_type=flow_type,
        )
        self.records.append(rec)
        logger.debug(
            "[CostTracker] %s | %s | in=%d out=%d | $%.5f",
            item_name, model_id, input_tokens, output_tokens, cost,
        )

    def save_report(self, label: str = "", flow_type: str = "Unknown") -> str:
        """Save cost report to common JSON file (token_log.json), append new records."""
        path = Path(LOG_DIR) / "token_log.json"
        os.makedirs(LOG_DIR, exist_ok=True)

        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        new_session = {
            "timestamp": ts,
            "screen_name": label or "unknown",
            "flow_type": flow_type,
            "total_requests": len(self.records),
            "total_input_tokens": self.total_input_tokens,
            "total_output_tokens": self.total_output_tokens,
            "total_tokens": self.total_tokens,
            "total_cost_usd": round(self.total_cost, 6),
            "records": [
                {
                    "item": r.item_name,
                    "model": r.model_id,
                    "stage": r.stage,
                    "input_tokens": r.input_tokens,
                    "output_tokens": r.output_tokens,
                    "cost_usd": round(r.cost_usd, 6),
                    "flow_type": r.flow_type
                } for r in self.records
            ]
        }

        with _REPORT_LOCK:
            history = []
            if path.exists():
                try:
                    with open(path, "r", encoding="utf-8") as f:
                        history = json.load(f)
                        if not isinstance(history, list):
                            history = [history] # Convert old format
                except Exception as exc:
                    if not _move_aside(path, exc):
                        return ""
                    history = []

            history.append(new_session)

            tmp_path = path.with_name(f"{path.name}.{os.getpid()}.tmp")
            with open(tmp_path, "w", encoding="utf-8") as f:
                json.dump(history, f, ensure_ascii=False, indent=2)
            os.replace(tmp_path, path)
            
        logger.info("[CostTracker] JSON report appended → %s", path)
        return str(path)

    def save_excel_report(self, label: str = "", flow_type: str = "Unknown") -> str:
        """Export cost report to common Excel file (token_log.xlsx), append new row."""
        try:
            import openpyxl
            from openpyxl.styles import PatternFill, Border, Side
        except ImportError:
            logger.warning("[CostTracker] openpyxl not installed — Excel export skipped.")
            return ""

        os.makedirs(LOG_DIR, exist_ok=True)
        path = Path(LOG_DIR) / "token_log.xlsx"
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")

        with _REPORT_LOCK:
            wb = None
            if path.exists():
                try:
                    wb = openpyxl.load_workbook(str(path))
                    ws = wb.active
                except Exception as exc:
                    if not _move_aside(path, exc):
                        return ""
                    wb = None
            if wb is None:
                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = "Token Usage History"
                self._create_excel_header(ws)

            # ── Styles ────────────────────────────────────────────────────────
            thin_border = Border(left=Side(style='thin'), right=Side(style='thin'),
                                top=Side(style='thin'), bottom=Side(style='thin'))
            alt_fill = PatternFill("solid", fgColor="F3F4F6")

            # ── Append Records ────────────────────────────────────────────────
            for i, r in enumerate(self.records, 1):
                total_tok = r.input_tokens + r.output_tokens
                # Format: ID, Date, Flow, Screen, Item, Model, Stage, In, Out, Total, Cost
                row = [
                    ws.max_row,
                    ts,
                    r.flow_type,
                    label,
                    r.item_name,
                    r.model_id,
                    r.stage,
                    r.input_tokens,
                    r.output_tokens,
                    total_tok,
                    round(r.cost_usd, 6)
                ]
                ws.append(row)
            
                # Apply basic styling to the new row
                for cell in ws[ws.max_row]:
                    cell.border = thin_border
                    if (ws.max_row - 1) % 2 == 0:
                        cell.fill = alt_fill

            try:
                wb.save(str(path))
                logger.info("[CostTracker] Excel report appended → %s", path)
            except Exception as e:
                logger.error("[CostTracker] Failed to save Excel (maybe file is open?): %s", e)

        return str(path)

    def _create_excel_header(self, ws):
        from openpyxl.styles import Font, PatternFill, Alignment
        headers = ["ID", "Timestamp", "Flow", "Screen", "Item", "Model", "Stage", "Input", "Output", "Total", "Cost (USD)"]
        ws.append(headers)
        header_fill = PatternFill("solid", fgColor="1E3A5F")
        header_font = Font(bold=True, color="FFFFFF")
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
        
        # Set widths
        widths = [5, 18, 12, 25, 30, 20, 20, 10, 10, 10, 12]
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w


# ── Singleton ──────────────────────────────────────────────────────────────
_tracker: CostTracker | None = None


def get_tracker() -> CostTracker:
    """Returns the singleton CostTracker for the current session."""
    global _tracker
    if _tracker is None:
        _tracker = CostTracker()
    return _tracker
